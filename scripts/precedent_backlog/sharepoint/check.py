"""Content-check each queued SharePoint candidate: download it, read the
opening, hand the file back to the cloud, and ask the model what it is.

Resumable: results append to data/checked.jsonl, keyed by path.
    python3 check.py [limit]
"""
import hashlib, io, json, os, re, shutil, subprocess, sys, threading, time, zipfile
from concurrent.futures import ThreadPoolExecutor

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
BACKLOG = os.path.dirname(HERE)
sys.path.insert(0, BACKLOG)
from dotenv import load_dotenv  # noqa: E402
load_dotenv(os.path.join(BACKLOG, ".env"))
import anthropic, classify  # noqa: E402

OUT = f"{DATA}/checked.jsonl"
EVICT = os.path.join(HERE, "evict")
MIN_FREE = 1_000_000_000
types = json.load(open(f"{DATA}/types.json"))
type_ids = {t["id"] for t in types}
listing = "\n".join(f"- {t['id']}: {t['name']} ({t['category']})" for t in types)
client = anthropic.Anthropic()
lock = threading.Lock()

SYSTEM = """You are curating a Pakistani law firm's precedent library: documents its lawyers will draft new work from. You get one file's name, its folder path, and the opening of its text, plus the library's document types.

Decide whether it belongs in the library and where.

Keep (keep=true) a complete instrument or document of a kind the firm drafts or negotiates: agreements, deeds, guarantees, undertakings, amendments, legal opinions, the firm's fee/engagement proposals, formal letters and notices, board resolutions, powers of attorney, certificates/consents/NOCs the firm drafts, due diligence reports, research memos, RFP documents, meeting minutes.

Reject (keep=false): working papers (comparisons, redlines, issue lists, trackers, questions lists, briefs, internal notes on another document); evidence and filing papers that are not drafting models (proof an account was opened, registry extracts, certificates of incorporation, bank statements, CNIC copies, invoices, receipts, court orders and judgments); cover pages, placeholders ("intentionally left blank"), a schedule or annex on its own, or a fragment; files with no readable text.

For a kept file, choose the single best existing type and return its id. Only if no existing type is a reasonable home, propose a new type: a short, standard name (e.g. "Share Purchase Agreement", "Accounts Agreement", "Security Trust Deed", "Guarantee", "Assignment Deed") and one of the existing categories. Classify by what the document IS, not by documents it mentions: a legal opinion on a concession agreement is a Legal Opinion; an amendment to an EPC contract is an Amendment Agreement.

Also give identity: "<instrument> | <project or deal> | <main parties>", in plain words, so copies of the same document across folders can be matched.

Respond with ONLY a JSON object:
{"keep": true|false, "type_id": "<uuid or null>", "new_type_name": "<string or null>", "new_type_category": "<string or null>", "identity": "<string>", "reason": "<one short sentence>"}
type_id and new_type_name are mutually exclusive; both null when keep is false."""


def snippet_of(data: bytes, ext: str) -> tuple[str, int]:
    """Opening text, and a rough measure of size (chars for docx, pages for pdf)."""
    if ext == "docx":
        import docx
        try:
            d = docx.Document(io.BytesIO(data))
        except KeyError:
            # python-docx trips over a dangling customXML part; the body
            # itself is fine, so read the text straight from the XML.
            xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf8", "ignore")
            full = re.sub(r"<[^>]+>", "", re.sub(r"<w:p[ >]", "\n<", xml))
            full = re.sub(r"[ \t]+", " ", full)
            return full[:3500], len(full)
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables[:3]:
            for row in t.rows[:8]:
                paras.append("\t".join(c.text for c in row.cells))
        full = "\n".join(paras)
        return full[:3500], len(full)
    import pdfplumber
    text = ""
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n = len(pdf.pages)
        for page in pdf.pages[:3]:
            text += (page.extract_text() or "") + "\n"
            if len(text) > 3500:
                break
    return text[:3500], n


def ask(row, snippet):
    msg = (
        f"Filename: {row['filename']}\nFolder: {os.path.dirname(row['rel'])}\n\nDocument types:\n{listing}\n\n"
        f"Opening text:\n{snippet[:3000] or '(no text extracted)'}"
    )
    r = client.messages.create(model=classify.MODEL, max_tokens=1500, system=SYSTEM,
                               messages=[{"role": "user", "content": msg}])
    text = next((b.text for b in r.content if b.type == "text"), "{}")
    data = json.loads(re.sub(r"^```json\s*|```\s*$", "", text.strip()))
    if data.get("type_id") and data["type_id"] not in type_ids:
        data["reason"] = f"(unknown type id returned) {data.get('reason', '')}"
        data["type_id"] = None
    return data


def wait_for_space(need):
    """Never let a download push the Mac's disk to full."""
    while shutil.disk_usage(os.path.expanduser("~")).free < MIN_FREE + need:
        print(f"  (low disk, waiting: {shutil.disk_usage(os.path.expanduser('~')).free // 1_000_000} MB free)", flush=True)
        time.sleep(60)


def work(row):
    res = {"path": row["path"]}
    try:
        wait_for_space(row["size"])
        with open(row["path"], "rb") as f:
            data = f.read()
        subprocess.run([EVICT, row["path"]], capture_output=True)
        res["sha1"] = hashlib.sha1(data).hexdigest()
        snip, size = snippet_of(data, row["ext"])
        res["measure"] = size
        res["snippet_chars"] = len(snip.strip())
        if len(snip.strip()) < 200:
            res.update(keep=None, reason="no readable text (scan?)", head=snip[:300])
        else:
            res.update(ask(row, snip))
            res["head"] = re.sub(r"\s+", " ", snip)[:600]
    except Exception as e:
        res.update(keep=None, error=f"{type(e).__name__}: {e}")
        subprocess.run([EVICT, row["path"]], capture_output=True)
    with lock:
        with open(OUT, "a") as f:
            f.write(json.dumps(res) + "\n")
    return res


if __name__ == "__main__":
    queue = json.load(open(f"{DATA}/queue.json"))
    done = set()
    if os.path.exists(OUT):
        for line in open(OUT):
            r = json.loads(line)
            if "error" not in r:
                done.add(r["path"])
    todo = [r for r in queue if r["path"] not in done]
    todo = todo[: int(sys.argv[1])] if len(sys.argv) > 1 else todo
    print(f"{len(todo)} to check ({len(done)} done already)", flush=True)
    with ThreadPoolExecutor(10) as ex:
        for n, _ in enumerate(ex.map(work, todo), 1):
            if n % 25 == 0:
                print(f"{n}/{len(todo)}", flush=True)
    print("finished", flush=True)
