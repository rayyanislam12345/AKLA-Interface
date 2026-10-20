#!/usr/bin/env python3
"""Check a cut pack and write its REVIEW.txt for the owner.

  verify_pack.py <config.json> <pack-name> <repo-root>

Reads <out_dir>/<kind>-pack.docx, kinds.json and report.json (written by
build_pack.py build). Checks that the docx unzips and every XML part parses,
that python-docx opens it (when installed), runs the Hub's own reader
(chat-service/docxAgent.js inspectDocx) and format fingerprint
(chat-service/formatFingerprint.js fingerprint) on it read-only, lists the
capitalised names and numbers still in a FORM pack, and writes REVIEW.txt.
"""
import collections
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_pack as bp  # noqa: E402

KIND_KEYS = {"pPr", "rPr", "styleId", "numId", "ilvl", "label"}
TABLE_KEYS = {"tblPr", "tblGrid", "header", "body", "label"}
ROW_KEYS = {"trPr", "tcPr", "pPr", "rPr"}
FOOT_KEYS = {"pPr", "rPr", "refRPr"}


def check_kinds_shape(kj):
    """The contract with the engine: exactly these keys and types."""
    errs = []
    if set(kj) != {"kinds", "tables", "footnote"}:
        errs.append("top-level keys are %s" % sorted(kj))
    for name, k in kj.get("kinds", {}).items():
        if set(k) - {"restart"} != KIND_KEYS:
            errs.append("kind %s keys %s" % (name, sorted(k)))
        if "restart" in k and not isinstance(k["restart"], bool):
            errs.append("kind %s restart not true/false" % name)
        if k.get("pPr") and not (k["pPr"].startswith("<w:pPr") and k["pPr"].endswith("</w:pPr>") or k["pPr"].endswith("/>")):
            errs.append("kind %s pPr is not a w:pPr element" % name)
        if k.get("rPr") and not k["rPr"].startswith("<w:rPr"):
            errs.append("kind %s rPr is not a w:rPr element" % name)
        for f in ("numId", "ilvl"):
            if k.get(f) is not None and not isinstance(k[f], int):
                errs.append("kind %s %s not int/null" % (name, f))
        if k.get("styleId") is not None and not isinstance(k["styleId"], str):
            errs.append("kind %s styleId not str/null" % name)
        for f in ("pPr", "rPr"):
            if k.get(f):
                try:
                    etree.fromstring(_ns_wrap(k[f]))
                except etree.XMLSyntaxError as e:
                    errs.append("kind %s %s does not parse: %s" % (name, f, e))
    for name, t in kj.get("tables", {}).items():
        if set(t) != TABLE_KEYS:
            errs.append("table %s keys %s" % (name, sorted(t)))
        for row in ("header", "body"):
            if set(t.get(row, {})) != ROW_KEYS:
                errs.append("table %s %s keys %s" % (name, row, sorted(t.get(row, {}))))
    fn = kj.get("footnote")
    if fn is not None and set(fn) != FOOT_KEYS:
        errs.append("footnote keys %s" % sorted(fn))
    return errs


def _ns_wrap(xml):
    decl = ('xmlns:w="%s" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
            'xmlns:r="%s" xmlns:mc="%s" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"') % (bp.W_NS, bp.R_NS, bp.MC_NS)
    return "<x %s>%s</x>" % (decl, xml)


def open_checks(pack):
    out = {"zip": "ok", "xml_parts": 0, "xml_errors": [], "python_docx": None}
    try:
        with zipfile.ZipFile(pack) as zf:
            bad = zf.testzip()
            if bad:
                out["zip"] = "bad member %s" % bad
            for n in zf.namelist():
                if n.endswith(".xml") or n.endswith(".rels"):
                    out["xml_parts"] += 1
                    try:
                        etree.fromstring(zf.read(n), bp.PARSER)
                    except etree.XMLSyntaxError as e:
                        out["xml_errors"].append("%s: %s" % (n, e))
    except zipfile.BadZipFile as e:
        out["zip"] = "not a zip: %s" % e
    try:
        import docx  # noqa
        d = docx.Document(pack)
        out["python_docx"] = "opens: %d paragraphs, %d tables, %d sections" % (len(d.paragraphs), len(d.tables), len(d.sections))
    except ImportError:
        out["python_docx"] = "not installed"
    except Exception as e:  # noqa
        out["python_docx"] = "FAILED: %s" % e
    return out


NODE = r"""
const [repo, pack] = process.argv.slice(-2);
const { readFileSync } = await import('node:fs');
const { pathToFileURL } = await import('node:url');
const agent = await import(pathToFileURL(repo + '/chat-service/docxAgent.js').href);
const fpm = await import(pathToFileURL(repo + '/chat-service/formatFingerprint.js').href);
const bytes = new Uint8Array(readFileSync(pack));
const out = {};
try {
  const v = agent.inspectDocx(bytes);
  out.inspect = { paragraphCount: v.paragraphCount, shown: v.shown, partial: v.partial,
                  placeholders: v.quality.placeholders.length, errors: v.quality.errors,
                  firstLines: v.listing.split('\n').slice(0, 12) };
} catch (e) { out.inspect = { error: String(e && e.message || e) }; }
try { out.fingerprint = fpm.fingerprint(bytes); } catch (e) { out.fingerprint = { error: String(e && e.message || e) }; }
console.log(JSON.stringify(out));
"""


def hub_checks(repo, pack):
    try:
        r = subprocess.run(["node", "--input-type=module", "-e", NODE, "--", repo, pack],
                           capture_output=True, text=True, timeout=180, cwd=os.path.join(repo, "chat-service"))
        if r.returncode != 0:
            return {"error": (r.stderr or r.stdout)[-1500:]}
        return json.loads(r.stdout.strip().splitlines()[-1])
    except Exception as e:  # noqa
        return {"error": str(e)}


def numbering_check(pack):
    """Every number in the pack read three ways: the builder's counter with
    one counter per list, with counters shared by a list's nums, and
    LibreOffice's own rendering (when installed). Also any typed number
    still at the start of a paragraph."""
    pkg = bp.Package(pack)
    styles = bp.style_map(pkg)
    nb = bp.Numbering(pkg, styles)
    body = pkg.xml(pkg.main_part()).find(bp.w("body"))
    paras = bp.body_paragraphs(body)
    toc = bp._field_state_paragraphs(body)
    ms = {m: bp.numbering_labels(body, styles, nb, m) for m in bp.MODELS}
    m1 = ms["own"]
    out = {"numbered": len(m1), "models_disagree": [], "typed_left": [], "libreoffice": None}
    for p in paras:
        if p in m1:
            if len({ms[m].get(p) for m in bp.MODELS}) > 1:
                out["models_disagree"].append(dict({"text": bp.para_text(p)[:60]}, **{m: ms[m].get(p) for m in bp.MODELS}))
            continue
        ppr = p.find(bp.w("pPr"))
        sid = bp.wa(ppr.find(bp.w("pStyle")), "val") if ppr is not None else None
        t = bp.para_text(p)
        if bp.TYPED_NUM.match(t) and p not in toc and not (sid or "").upper().startswith("TOC") and t[bp.TYPED_NUM.match(t).end():].strip():
            out["typed_left"].append(t[:70])
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        out["libreoffice"] = "not installed"
        return out
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "pack.docx")
        shutil.copy(pack, src)
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "txt:Text (encoded):UTF8", "--outdir", tmp, src],
                           capture_output=True, timeout=300)
            with open(os.path.join(tmp, "pack.txt"), encoding="utf-8-sig") as f:
                lines = f.read().split("\n")
        except Exception as e:  # noqa
            out["libreoffice"] = "could not render: %s" % e
            return out
    # walk the pack's paragraphs and LibreOffice's lines together: a line is
    # a paragraph's text, after its number when it has one
    norm = lambda t: re.sub(r"\s+", " ", t.replace("\t", " ")).strip()
    lines = [norm(x) for x in lines]
    at, same, differ, unfound = 0, 0, [], 0
    for p in paras:
        text = norm(bp.para_text(p))
        if len(re.sub(r"\[●\]|\W", "", text)) < 6:
            # too little of its own to find it in the rendered text ("[●]")
            if p in m1:
                unfound += 1
            continue
        hit = None
        for k in range(at, min(len(lines), at + 60)):
            ln = lines[k]
            numbered = p in m1 and (m1[p] or "").strip()
            if (not numbered and ln == text) or (numbered and ln.endswith(" " + text) and len(ln) - len(text) <= 16):
                hit = k
                break
        if hit is None:
            if p in m1:
                unfound += 1
            continue
        at = hit + 1
        if p not in m1:
            continue
        label = lines[hit][:len(lines[hit]) - len(text)].strip()
        readings = sorted({(ms[m].get(p) or "").strip() for m in bp.MODELS})
        bullet = lambda t: len(t) == 1 and not t.isalnum()
        if label in readings or (bullet(label) and all(bullet(r) for r in readings)):
            same += 1
        else:
            differ.append({"text": text[:40], "builder": " or ".join(readings), "libreoffice": label})
    out["libreoffice"] = {"agree": same, "differ": differ, "not_found": unfound}
    return out


# a project may be federal or under any province: no pack names a
# government, a province, a provincial body or a province-only law
FORBIDDEN = re.compile(r"Government of (?!\[)[A-Z][\w-]*|\b(?:GoS|GOS|GoP|GOP|GoPb|GoKP|GoB)\b|Sindh|Punjab|Khyber\s*Pakhtunkhwa|\bKPK\b|"
                       r"Balochistan|Baluchistan|Islamabad|Karachi|Lahore|Peshawar|Quetta|Hyderabad|Multan|Faisalabad|Rawalpindi|"
                       r"Province of (?!\[)[A-Z][\w-]*|Provincial Government|Revenue Board|\bSRB\b|\bKPRA\b|\bPRA\b|\bBRA\b|"
                       r"Public[- ]Private Partnership (?:Authority|Act)|\bP3A\b|PPP (?:Authority|Act)|Procurement Rules|"
                       r"Procurement Regulatory Authority|\b(?:S|K|B)?PPRA\b|Local Government|Housing (?:&|and) Town")
# laws and bodies that apply whatever the project: kept, and listed
NATIONWIDE = re.compile(r"Companies Act(?:,? \d{4})?|\bSECP\b|Securities and Exchange Commission of Pakistan|Arbitration Act(?:,? 1940)?|"
                        r"Contract Act(?:,? 1872)?|Federal Board of Revenue(?: \(or any other authority\))?|\bFBR\b|State Bank of Pakistan|"
                        r"Income Tax Ordinance(?:,? 2001)?|Sales Tax Act(?:,? 1990)?|Stamp Act(?:,? 1899)?|Registration Act(?:,? 1908)?|"
                        r"Code of Civil Procedure(?:,? 1908)?|Limitation Act(?:,? 1908)?")
NEUTRAL_REASON = re.compile(r"government|provinc|\bGoS\b|sales tax authority", re.I)


def government_check(pack, kj):
    found, kept = [], collections.Counter()
    with zipfile.ZipFile(pack) as zf:
        for n in zf.namelist():
            if not bp.STORY.match(n):
                continue
            root = etree.fromstring(zf.read(n), bp.PARSER)
            for p in root.iter(bp.w("p")):
                t = bp.para_text(p)
                for m in FORBIDDEN.finditer(t):
                    found.append("%s: …%s…" % (n.replace("word/", ""), t[max(0, m.start() - 40):m.end() + 40].replace("\t", " ")))
                for m in NATIONWIDE.finditer(t):
                    kept[m.group(0)] += 1
    for name, k in list(kj.get("kinds", {}).items()) + list(kj.get("tables", {}).items()):
        for m in FORBIDDEN.finditer(k.get("label", "")):
            found.append("kinds.json %s label: %s" % (name, k["label"]))
    return {"found": found, "kept": kept.most_common()}


LEAD = {"The", "This", "That", "These", "Those", "Such", "Each", "Any", "All", "No", "Every", "Other", "In", "On", "At", "By",
        "For", "If", "Upon", "Where", "When", "Subject", "Unless", "Notwithstanding", "Without", "Save", "Pursuant", "Under",
        "With", "From", "To", "As", "A", "An", "Its", "Their", "Both", "Either", "Neither", "Following", "Prior", "After",
        "Before", "Within", "During", "Except", "Provided", "Further", "Accordingly", "Now", "Whereas", "And", "Or", "Of",
        "It", "Is", "Be", "Shall", "May", "Not", "Has", "Have"}
CAPSEQ = re.compile(r"\b[A-Z][\w&'’.\-]*(?:[ ]+(?:of|and|&|the|for|de|al|ul|bin|e|du)?[ ]*[A-Z][\w&'’.\-]*)+")
ACRO = re.compile(r"\b[A-Z][A-Za-z]*[A-Z][A-Za-z]*\b")
NUM = re.compile(r"(?<![\w.])\d[\d,./-]*\d?(?![\w])")
REFWORD = re.compile(r"(Clause|Clauses|Section|Sections|Article|Articles|Schedule|Annex|Annexure|Part|paragraph|Paragraph|Rule|Regulation|Item|item|sub-clause|Sub-Clause|No\.|Page|page|of|Appendix|Exhibit|Form|Table|Chapter|and|or|to)\s*$|(?:clause|clauses|section|sections|article|articles|schedule|regulation|rule)s?\s+[\d.()a-z, ]*$", re.I)
WORDNUM = re.compile(r"\b(?:one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|twenty|thirty|forty|fifty|sixty|seventy|eighty|ninety|hundred|thousand)[\s-]*\($", re.I)
STATUTE = re.compile(r"(Act|Ordinance|Rules|Regulations|Code|Order|Law|Laws),?\s*$")


def leftovers(pack):
    """Capitalised names and numbers still in the pack outside [●]."""
    texts = []
    with zipfile.ZipFile(pack) as zf:
        for n in zf.namelist():
            if bp.STORY.match(n):
                root = etree.fromstring(zf.read(n), bp.PARSER)
                for p in root.iter(bp.w("p")):
                    t = bp.para_text(p)
                    if t.strip():
                        texts.append(t)
    full = "\n".join(texts)
    defined = set()
    for m in re.finditer(r"[“\"‘]([^”\"’\n]{2,70})[”\"’]", full):
        defined.add(m.group(1).strip())
    for m in re.finditer(r"\((?:the|this|each|together,? the|collectively,? the)\s+([A-Z][^()\n]{1,60})\)", full):
        defined.add(m.group(1).strip())
    for m in re.finditer(r"^\s*([A-Z][\w &'’\-]{1,60}?)\s+(?:means|has the meaning|shall mean|shall have the meaning)", full, re.M):
        defined.add(m.group(1).strip())
    names = collections.Counter()
    for m in CAPSEQ.finditer(full):
        words = m.group(0).split()
        while words and words[0] in LEAD:
            words = words[1:]
        if len(words) < 2:
            continue
        names[" ".join(words).strip(".,;:")] += 1
    acr = collections.Counter(m.group(0) for m in ACRO.finditer(full) if len(m.group(0)) >= 2)
    nums = []
    for t in texts:
        for m in NUM.finditer(t):
            s = m.group(0)
            before = t[max(0, m.start() - 25):m.start()]
            if re.fullmatch(r"\d{1,2}(\.\d{1,2})*\.?", s) and (REFWORD.search(before) or m.start() < 3):
                continue
            if re.fullmatch(r"(19|20)\d{2}", s) and STATUTE.search(before):
                continue
            if REFWORD.search(before):
                continue
            if WORDNUM.search(before) and t[m.end():m.end() + 1] == ")":
                continue   # "thirty (30)": a time period or count, listed with the kept time periods
            nums.append("%s  …%s[%s]%s…" % (s, before[-22:], s, t[m.end():m.end() + 22]))
    def is_defined(n):
        return n in defined or any(n.startswith(d) and len(d) > 3 for d in defined)
    check = [(n, c) for n, c in names.most_common() if not is_defined(n)]
    deflike = [(n, c) for n, c in names.most_common() if is_defined(n)]
    return {"check_names": check, "defined_or_headings": deflike, "acronyms": acr.most_common(), "numbers": nums}


def write_review(cfg, out_dir, repo):
    kind = cfg["kind"]
    pack = os.path.join(out_dir, "%s-pack.docx" % kind)
    with open(os.path.join(out_dir, "report.json"), encoding="utf-8") as f:
        rep = json.load(f)
    with open(os.path.join(out_dir, "kinds.json"), encoding="utf-8") as f:
        kj = json.load(f)
    oc = open_checks(pack)
    hub = hub_checks(repo, pack)
    shape_errs = check_kinds_shape(kj)
    lo = leftovers(pack)
    numck = numbering_check(pack)
    L = []
    add = L.append
    title = cfg.get("title", kind)
    add("%s PACK  (%s pack)" % (title.upper(), rep["mode"].upper()))
    add("=" * 72)
    add("Cut from: %s" % rep["base"])
    if cfg.get("base_note"):
        add(cfg["base_note"])
    add("Files here: %s-pack.docx (the pack), kinds.json (the paragraph kinds the drafting engine" % kind)
    add("writes with), report.json (the full machine report), REVIEW.txt (this note).")
    add("")
    add("WHY %s" % rep["mode"].upper())
    for line in cfg.get("why", []):
        add("  " + line)
    add("")
    add("1. WHAT WAS KEPT")
    if rep["mode"] == "shell":
        sh = rep["shell"]
        add("  The base had %d top-level paragraphs and tables. %d were deleted; %d furniture items and" % (sh["blocks_before"], sh["deleted"], sh["furniture_kept"]))
        add("  %d example paragraphs/tables were kept (example text replaced by [●])." % sh["examples"])
        for t in sh.get("toc") or []:
            add("  Contents: " + t)
    else:
        add("  The whole text of the base, with revisions accepted, comments removed and the deal facts")
        add("  below replaced by [●].")
    for line in cfg.get("kept", []):
        add("  - " + line)
    if rep.get("letterhead"):
        add("  Letterhead: " + "; ".join(rep["letterhead"]))
    add("")
    add("2. WHAT WAS BLANKED (every replacement; 'part' is where in the file)")
    if not rep["blankings"]:
        add("  nothing")
    for b in rep["blankings"]:
        txt = b["text"].replace("\n", " / ").replace("\t", " ")
        if len(txt) > 160:
            txt = txt[:157] + "..."
        where = b["part"].replace("word/", "") + (" " + b["where"] if b["where"] else "")
        add("  - [%s] %s  ->  [●]   (%s)" % (where, txt, b["reason"]))
    if rep.get("time_periods_kept"):
        add("  Kept on purpose (time periods are settled terms, not deal facts): " + "; ".join(rep["time_periods_kept"][:30]))
    if rep.get("properties_blanked"):
        add("  File properties: " + "; ".join(rep["properties_blanked"]))
    gov = government_check(pack, kj)
    add("  GOVERNMENT AND PROVINCE:")
    if gov["found"]:
        add("  !! Still named in the pack (a project may be federal or under any province):")
        for f in gov["found"][:40]:
            add("    - " + f)
    else:
        add("  No government, province or province-specific law is named; the government party is the Government.")
    subs = collections.Counter()
    for b in rep["blankings"]:
        if NEUTRAL_REASON.search(b["reason"]) or FORBIDDEN.search(b["text"]):
            to = b["reason"].split(" -> ", 1)[1] if " -> " in b["reason"] else "[●]"
            subs["'%s' -> '%s'" % (b["text"].replace("\n", " ").replace("\t", " ")[:90], to)] += 1
    if subs:
        add("  Substitutions made for this (each with how many times):")
        for k, n in subs.most_common():
            add("    - %s  (x%d)" % (k, n))
    else:
        add("  Substitutions made for this: none needed.")
    if gov["kept"]:
        add("  Kept as written (they apply nationwide whatever the project): " + "; ".join("%s (x%d)" % kv for kv in gov["kept"]))
    add("")
    add("3. KINDS AVAILABLE (kinds.json)")
    for name, k in kj["kinds"].items():
        num = " numbering %s level %s" % (k["numId"], k["ilvl"]) if k["numId"] is not None else ""
        add("  %-22s %s  [style %s%s]" % (name, k["label"], k["styleId"] or "none", num))
    for name, t in kj["tables"].items():
        add("  %-22s table: %s" % (name, t["label"]))
    add("  %-22s %s" % ("footnote", "yes, from the pack's footnote" if kj["footnote"] else "none in this pack"))
    if shape_errs:
        add("  !! kinds.json shape problems: " + "; ".join(shape_errs))
    add("")
    add("4. NUMBERING")
    an = rep.get("auto_number")
    if numck["typed_left"]:
        add("  Typed numbers still in the pack (not Word's numbering):")
        for t in numck["typed_left"]:
            add("    - " + t.replace("\t", " -> "))
    else:
        add("  All numbering is Word's automatic numbering (%d numbered paragraphs; no clause or limb number is typed)." % numck["numbered"])
    if an:
        conv = an["converted"]
        add("  Typed numbers converted to Word list items: %d" % len(conv))
        for lk, n in sorted(an["lists_used"].items(), key=lambda x: -x[1]):
            add("    %3d onto %s" % (n, lk))
        if conv:
            add("    (one by one: " + "; ".join("%s->num %s/level %s" % (c["typed"], c["numId"], c["ilvl"] + 1) for c in conv[:400]) + ")")
        ind = an.get("indents") or {}
        if ind:
            add("  Indents: %d took the list level's own indent (their own was the same and was dropped); %d kept their own indent,"
                % (ind.get("dropped", 0) + ind.get("level", 0), ind.get("kept", 0) + ind.get("kept (from its style)", 0)))
            add("    written out in full where it differs from the list level, so nothing moves on the page.")
        na = an["numbering_added"]
        if na["nums"] or na["abstracts"] or na.get("levels_changed"):
            add("  numbering.xml additions (the pack's numbering becomes the firm's standard for this pack):")
            by = collections.Counter(n["abstractNumId"] for n in na["nums"])
            add("    %d list starts (a w:num on an existing list, with a start override where a typed list starts again or"
                " starts part-way): %s" % (len(na["nums"]), ", ".join("list %s x%d" % kv for kv in sorted(by.items()))))
            for a in na["abstracts"]:
                add("    new list %s: %s" % (a["abstractNumId"], a["why"]))
            for a in na.get("levels_changed", []):
                add("    list %s level %s, used by no paragraph, changed from %s to %s: %s" % (a["abstractNumId"], a["level"], a["was"], a["now"], a["why"]))
        else:
            add("  numbering.xml: nothing added.")
        for a in an["anomalies_kept"]:
            add("  !! Kept as the base has it: %s  ['%s']" % (a["why"], a["text"][:60].replace("\t", " ")))
        for t in an["left_typed"]:
            add("  !! Left typed: %s" % (t if isinstance(t, str) else "%s ['%s']" % (t["why"], t["text"][:60])))
        for t in an["not_numbers"]:
            add("  Not a list number, left as text: '%s'" % t["text"][:60])
        sc = an.get("sequence_check") or {}
        add("  Sequence check (builder): %d numbered paragraphs read exactly as before, typed or not, under each of three readings"
            " of how Word counts lists (%s); %d differences." % (sc.get("paragraphs_checked", 0), "; ".join(sc.get("models", [])), len(sc.get("mismatches", []))))
    else:
        add("  Automatic numbering was not run for this pack.")
    lon = numck["libreoffice"]
    if isinstance(lon, dict):
        add("  LibreOffice rendering: %d numbers agree with the builder's count, %d differ, %d not checked (text too short or too"
            " common to find in LibreOffice's text, e.g. '[●]')."
            % (lon["agree"], len(lon["differ"]), lon["not_found"]))
        for d in lon["differ"][:20]:
            add("    !! '%s': builder %s, LibreOffice %s" % (d["text"], d["builder"], d["libreoffice"]))
    else:
        add("  LibreOffice rendering: %s" % lon)
    base_amb = {d["text"] for d in (an or {}).get("base_ambiguous", [])}
    own = [d for d in numck["models_disagree"] if d["text"] in base_amb]
    numck["models_disagree"] = [d for d in numck["models_disagree"] if d["text"] not in base_amb]
    if own:
        add("  The base's own numbering (not typed, not changed here) reads differently depending on whether Word shares a list's")
        add("  counters across its nums at %d places (the readings: %s). Check these in Word" % (
            len(own), "; ".join("'%s' %s" % (d["text"][:30], " / ".join(sorted({str(d[m]) for m in bp.MODELS}))) for d in own[:12])))
    if numck["models_disagree"]:
        add("  !! %d numbers read differently depending on whether Word shares a list's counters across its nums: %s"
            % (len(numck["models_disagree"]), "; ".join("'%s' %s" % (d["text"][:30], " / ".join(sorted({str(d[m]) for m in bp.MODELS}))) for d in numck["models_disagree"][:10])))
    if rep.get("kinds_run_on"):
        add("  Kinds whose list runs on through the document (restart: false; the engine never starts them again): %s"
            % ", ".join(rep["kinds_run_on"]))
    add("")
    add("5. PAGE SET-UP")
    for s in rep["sections"]:
        m = s["margins_twips"] or {}
        add("  Section %d: %s %s, margins top %s / right %s / bottom %s / left %s twips; %s" % (
            s["section"], s["page"], s["orientation"], m.get("top"), m.get("right"), m.get("bottom"), m.get("left"),
            ", ".join(s["headers_footers"]) or "no header/footer"))
    for c in rep["page_changes"]:
        add("  Changed: " + c)
    if not rep["page_changes"]:
        add("  Already A4 portrait; nothing changed.")
    for t in rep.get("wide_tables", []):
        add("  !! Table at block %d is %d twips wide against a %d twip text width: check it does not run into the margin." % (t["block"], t["table_width_twips"] + t["indent"], t["text_width_twips"]))
    add("")
    add("6. SHAPES, REVISIONS, CONTENTS")
    if rep["shapes"]:
        for s in rep["shapes"]:
            add("  - %s: %s%s%s%s" % (s["part"].replace("word/", ""), "text box" if s["text_box"] else ("picture" if s["picture"] else "drawing/shape"),
                                     (" reading '%s'" % s["text"][:80]) if s["text"] else "",
                                     (" in paragraph '%s'" % s["paragraph_text"][:60]) if s["paragraph_text"].strip() else "",
                                     "  <- MUST BE FILLED BY HAND" if s["must_fill"] else ""))
    else:
        add("  No shapes or text boxes.")
    add("  Revisions accepted: %s; comments removed: %d." % (
        sum(sum(v.values()) for v in rep["revisions_accepted"].values()) if rep["revisions_accepted"] else 0, rep["comments_removed"]))
    add("  Revisions left in the pack: %s" % (rep["revisions_left"] or "none"))
    add("  Contents (TOC) field: %s" % ("; ".join(rep["toc_fields"]) if rep["toc_fields"] else "none"))
    if rep.get("field_imbalance"):
        add("  !! Field begin/end marks do not balance (%s)." % rep["field_imbalance"])
    add("  Parts changed: %s" % ", ".join(rep["changed_parts"]))
    add("  Parts removed: %s" % (", ".join(rep["removed_parts"]) or "none"))
    add("  Carried byte-for-byte: %s" % ", ".join(p for p in rep["carried_byte_for_byte"] if not p.startswith("customXml")))
    add("")
    add("7. OPENS, AND WHAT THE HUB READS")
    add("  Zip: %s; %d XML parts parse%s." % (oc["zip"], oc["xml_parts"], "" if not oc["xml_errors"] else ("; ERRORS: " + "; ".join(oc["xml_errors"]))))
    add("  python-docx: %s" % oc["python_docx"])
    ins = hub.get("inspect", {}) if "error" not in hub else {"error": hub["error"]}
    if "error" in ins:
        add("  Hub reader (inspectDocx): FAILED: %s" % ins["error"][:400])
    else:
        add("  Hub reader (inspectDocx): %d paragraphs, %d placeholders, table errors: %s" % (ins["paragraphCount"], ins["placeholders"], ins["errors"] or "none"))
    fp = hub.get("fingerprint", {}) if "error" not in hub else {"error": hub["error"]}
    if "error" in fp:
        add("  Format fingerprint: FAILED: %s" % fp["error"][:400])
    else:
        b = fp["body"]
        add("  Format fingerprint: body font %s (%.1f%% of the text), size %s pt (%.1f%% of the text), %d characters." % (
            b["font"], 100 * (b["fontShare"] or 0), (b["sizeHalfPts"] or 0) / 2, 100 * (b["share"] or 0), b["chars"]))
        add("    fonts: %s; sizes (half-points): %s" % (b["fonts"], b["sizes"]))
        add("    sections: %s" % "; ".join("%sx%s%s" % (s["pgSz"].get("w") if isinstance(s["pgSz"], dict) else s["pgSz"],
                                                      s["pgSz"].get("h") if isinstance(s["pgSz"], dict) else "", " " + s["orient"] if s.get("orient") else "") for s in fp["sections"]))
        add("    heading styles: %s; table styles: %s; placeholder: %s" % (fp.get("headingStyles"), fp.get("tableStyles"), fp.get("placeholder")))
    add("")
    if rep["mode"] == "form":
        add("8. NAMES AND NUMBERS STILL IN THE PACK (check none is the old client's)")
        add("  Capitalised names that are not defined terms in the pack (count):")
        add("    " + "; ".join("%s (%d)" % nc for nc in lo["check_names"][:120]) if lo["check_names"] else "    none")
        add("  Acronyms and capitalised words (count):")
        add("    " + "; ".join("%s (%d)" % nc for nc in lo["acronyms"][:80]) if lo["acronyms"] else "    none")
        add("  Defined terms and headings (for completeness, count):")
        add("    " + "; ".join("%s (%d)" % nc for nc in lo["defined_or_headings"][:80]))
        add("  Numbers not part of a clause reference (with context):")
        for n in lo["numbers"][:80]:
            add("    " + n)
        if len(lo["numbers"]) > 80:
            add("    ... and %d more (see report.json)" % (len(lo["numbers"]) - 80))
        add("")
    add("%d. ANYTHING ODD" % (9 if rep["mode"] == "form" else 8))
    for line in cfg.get("odd", []) or ["nothing beyond the above"]:
        add("  - " + line)
    add("")
    add("%d. CHECK THESE BY EYE IN WORD" % (10 if rep["mode"] == "form" else 9))
    for i, line in enumerate(cfg.get("checks", []), 1):
        add("  %d. %s" % (i, line))
    add("")
    with open(os.path.join(out_dir, "REVIEW.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(L))
    rep["verification"] = {"open": oc, "hub": hub, "kinds_shape_errors": shape_errs, "leftovers": lo, "numbering": numck, "government": gov}
    with open(os.path.join(out_dir, "report.json"), "w", encoding="utf-8") as f:
        json.dump(rep, f, ensure_ascii=False, indent=2)
    return oc, hub, shape_errs, lo, numck, gov


def main(argv):
    if len(argv) < 4:
        print(__doc__)
        return 2
    with open(argv[1], encoding="utf-8") as f:
        cfgs = json.load(f)
    cfg = cfgs["packs"][argv[2]]
    out_dir = os.path.join(cfgs["out_root"], cfg["folder"])
    oc, hub, se, lo, numck, gov = write_review(cfg, out_dir, argv[3])
    ok = oc["zip"] == "ok" and not oc["xml_errors"] and not se and "error" not in hub and "error" not in hub.get("inspect", {}) and "error" not in hub.get("fingerprint", {})
    lo_r = numck["libreoffice"]
    ok = ok and not numck["typed_left"] and not numck["models_disagree"] and not (isinstance(lo_r, dict) and lo_r["differ"]) and not gov["found"]
    print("%s: %s" % (argv[2], "OK" if ok else "PROBLEMS"), json.dumps({"open": oc, "shape": se, "hub_err": hub.get("error"),
                                                                     "typed_left": numck["typed_left"][:5], "government_named": gov["found"][:5], "models_disagree": len(numck["models_disagree"]),
                                                                     "libreoffice": lo_r if not isinstance(lo_r, dict) else {"agree": lo_r["agree"], "differ": lo_r["differ"][:5], "not_found": lo_r["not_found"]}})[:1500])
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
