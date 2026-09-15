#!/usr/bin/env python3
"""Render a document into a Word file in AKLA house format.

The house format comes from the firm's Word Formatting and Shortcuts Guide
(June 17, 2026) and was measured against circulated firm documents - the
Artistic DISCOs Proposal (AM7), the Notes on Restrictions under the MRA
(OES2), the FESCO data-room Memorandum (FE-M01) and the Special Technology
Zone formatting exercise - with scripts/akla_style/analyze_docx.py. Where those
documents disagree, the firm chose (September 15, 2026):

  - section headings are numbered "1.", left-aligned, 13pt bold small caps,
    with a single rule beneath, in memos, notes, proposals and reports alike;
  - the navy banner with gold small caps is for the document title only;
  - body text sits one level below its heading: 1.1. under a section,
    1.1.1. under a sub-heading;
  - nothing is indented. Every number sits at the left margin and all text
    starts the same distance from it, at every level - headings, clauses,
    lists, quoted text and tables alike. Only the number shows the depth.

How it is built matters as much as how it looks. The first version typed
formatting onto every paragraph, so indents drifted and anything added later
in Word or by the AI had nothing to inherit. Here every paragraph takes a
named AKLA style, and the outline numbering belongs to those styles, so Word
renumbers and indents consistently and a paragraph inserted later picks up the
same format.

Input is a Markdown subset:

    # Title                 -> navy title banner, gold bold small caps
    lines before first ##   -> front matter (parties, date, status), unnumbered
    ## Section              -> AKLA Heading 1: "1." 13pt bold small caps, rule beneath
    ### Sub-section         -> AKLA Heading 2: "1.1." bold underlined small caps
    #### Divider            -> centred bold small caps with a rule beneath
    plain paragraph         -> AKLA Body 1 ("1.1.") or AKLA Body 2 ("1.1.1.")
    - item / 1. item        -> (a); indented two spaces -> (i); four -> A.
    > quoted text           -> italic, aligned with the clause text, justified
    ::: key figure          -> shaded centred box
    | a | b |               -> table, navy header row with gold small caps
    **bold**  *italic*      -> inline runs
    [[AKLA Comment: ...]]   -> native Word comment by "AKLA Comments"
    [^1] ... [^1]: note     -> the note becomes a comment where it is cited

Usage:
    python3 build_docx.py NOTES.md -o OUT.docx --doc-title "..." --doc-status "Draft" --doc-date "September 15, 2026"
"""
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

FONT = "Arial"
BODY_PT = 11
HEADING_PT = 13
TITLE_PT = 14
SMALL_PT = 8

NAVY = "002060"   # title banner, table header rows, the firm's name
GOLD = "FFC000"   # text on navy
RED = "C00000"    # confidentiality line in the running header
SHADE = "F2F2F2"  # key-figure box

TWIP = 1440                 # twentieths of a point in an inch
PAGE_W, PAGE_H = Mm(210), Mm(297)
MARGIN = Inches(1)
TEXT_TWIPS = int((PAGE_W - 2 * MARGIN) / 635)  # EMU -> twips

# Where the text of every numbered paragraph starts, at every level: wide
# enough for "10.10.10." in Arial 11 after a number set at the margin.
TEXT_AT = 1080              # twips (0.75")

OUTLINE = (7100, 7101)      # abstractNumId, numId for 1. / 1.1. / 1.1.1.
# Lists restart for each group, so each gets its own w:num; these abstract
# ids are the two indentation contexts a list can sit in.
LIST_UNDER_BODY1 = 7200     # (a) under a 1.1. paragraph
LIST_UNDER_BODY2 = 7300     # (a) under a 1.1.1. paragraph

FIRM_NAME = re.compile(r"(Ali Khan Law Associates)")

# Each of these element sequences is schema-mandated: append a child out of
# order and Word declares the file corrupt on open.
PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]
TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]
TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge",
    "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
    "w:tcFitText", "w:vAlign", "w:hideMark",
]
TRPR_ORDER = ["w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore",
              "w:wAfter", "w:cantSplit", "w:trHeight", "w:tblHeader",
              "w:tblCellSpacing", "w:jc", "w:hidden"]


# --------------------------------------------------------------------------
# XML helpers
# --------------------------------------------------------------------------

def _el(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el


def _insert(parent, element, sequence):
    tag = element.tag
    rank = sequence.index(next(s for s in sequence if qn(s) == tag))
    for existing in parent.findall(tag):
        parent.remove(existing)
    for i, child in enumerate(list(parent)):
        names = [qn(s) for s in sequence]
        if child.tag in names and names.index(child.tag) > rank:
            parent.insert(i, element)
            return element
    parent.append(element)
    return element


def ppr_insert(ppr, element):
    return _insert(ppr, element, PPR_ORDER)


def arial(rpr):
    """Arial for every script, and no theme font that would outrank it."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def bottom_rule(ppr, size=6, space=1):
    borders = OxmlElement("w:pBdr")
    borders.append(_el("w:bottom", **{"w:val": "single", "w:sz": size, "w:space": space, "w:color": "000000"}))
    ppr_insert(ppr, borders)


def numbering_ref(ppr, num_id, level):
    numpr = OxmlElement("w:numPr")
    numpr.append(_el("w:ilvl", **{"w:val": level}))
    numpr.append(_el("w:numId", **{"w:val": num_id}))
    ppr_insert(ppr, numpr)


def indent(ppr, left, hanging=None, right=None):
    attrs = {"w:left": left}
    if hanging is not None:
        attrs["w:hanging"] = hanging
    if right is not None:
        attrs["w:right"] = right
    ppr_insert(ppr, _el("w:ind", **attrs))


def cell_fill(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    _insert(tcpr, _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}), TCPR_ORDER)


def cell_margins(table, top, bottom, left=108, right=108):
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el(f"w:{edge}", **{"w:w": val, "w:type": "dxa"}))
    _insert(table._tbl.tblPr, mar, TBLPR_ORDER)


def table_borders(table, val="single"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        attrs = {"w:val": val}
        if val != "nil":
            attrs.update({"w:sz": 4, "w:space": 0, "w:color": "000000"})
        borders.append(_el(f"w:{edge}", **attrs))
    _insert(table._tbl.tblPr, borders, TBLPR_ORDER)


def table_width(table, twips, indent_twips=0):
    tblpr = table._tbl.tblPr
    _insert(tblpr, _el("w:tblW", **{"w:w": twips, "w:type": "dxa"}), TBLPR_ORDER)
    if indent_twips:
        _insert(tblpr, _el("w:tblInd", **{"w:w": indent_twips, "w:type": "dxa"}), TBLPR_ORDER)
    _insert(tblpr, _el("w:tblLayout", **{"w:type": "fixed"}), TBLPR_ORDER)


def add_field(paragraph, instruction, size_pt=SMALL_PT):
    """A Word field (PAGE, NUMPAGES) that updates when the file is opened."""
    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    run._r.append(placeholder)
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


# --------------------------------------------------------------------------
# numbering and styles
# --------------------------------------------------------------------------

def _level(ilvl, fmt, text, left, hanging, rpr=None):
    lvl = _el("w:lvl", **{"w:ilvl": ilvl})
    lvl.append(_el("w:start", **{"w:val": 1}))
    lvl.append(_el("w:numFmt", **{"w:val": fmt}))
    lvl.append(_el("w:lvlText", **{"w:val": text}))
    lvl.append(_el("w:lvlJc", **{"w:val": "left"}))
    ppr = OxmlElement("w:pPr")
    ppr.append(_el("w:ind", **{"w:left": left, "w:hanging": hanging}))
    lvl.append(ppr)
    if rpr is not None:
        lvl.append(rpr)
    return lvl


def build_numbering(doc):
    numbering = doc.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)

    # The outline. The firm does not indent: every level's number sits at the
    # margin and its text at TEXT_AT, so only the number shows the depth.
    outline = _el("w:abstractNum", **{"w:abstractNumId": OUTLINE[0]})
    outline.append(_el("w:multiLevelType", **{"w:val": "multilevel"}))
    plain = OxmlElement("w:rPr")
    plain.append(_el("w:u", **{"w:val": "none"}))
    outline.append(_level(0, "decimal", "%1.", TEXT_AT, TEXT_AT))
    outline.append(_level(1, "decimal", "%1.%2.", TEXT_AT, TEXT_AT, plain))
    outline.append(_level(2, "decimal", "%1.%2.%3.", TEXT_AT, TEXT_AT))
    for ilvl in range(3, 9):
        outline.append(_level(ilvl, "decimal", "%1.%2.%3." + "".join(f"%{i + 1}." for i in range(3, ilvl + 1)), TEXT_AT, TEXT_AT))
    numbering.append(outline)

    # Lists: (a) -> (i) -> A., the Guide's scheme, set the same way - number
    # at the margin, text at TEXT_AT.
    for abstract_id in (LIST_UNDER_BODY1, LIST_UNDER_BODY2):
        lists = _el("w:abstractNum", **{"w:abstractNumId": abstract_id})
        lists.append(_el("w:multiLevelType", **{"w:val": "multilevel"}))
        lists.append(_level(0, "lowerLetter", "(%1)", TEXT_AT, TEXT_AT))
        lists.append(_level(1, "lowerRoman", "(%2)", TEXT_AT, TEXT_AT))
        lists.append(_level(2, "upperLetter", "%3.", TEXT_AT, TEXT_AT))
        numbering.append(lists)

    num = _el("w:num", **{"w:numId": OUTLINE[1]})
    num.append(_el("w:abstractNumId", **{"w:val": OUTLINE[0]}))
    numbering.append(num)
    return numbering


class ListNumbers:
    """A fresh w:num per list, so every list starts again at (a)."""

    def __init__(self, numbering):
        self.numbering = numbering
        self.next_id = 7400

    def new(self, abstract_id):
        num_id = self.next_id
        self.next_id += 1
        num = _el("w:num", **{"w:numId": num_id})
        num.append(_el("w:abstractNumId", **{"w:val": abstract_id}))
        for ilvl in range(3):
            override = _el("w:lvlOverride", **{"w:ilvl": ilvl})
            override.append(_el("w:startOverride", **{"w:val": 1}))
            num.append(override)
        self.numbering.append(num)
        return num_id


def style(doc, name, base="Normal", size=None, bold=None, italic=None, small_caps=None,
          underline=None, color=None, align=None, before=None, after=None,
          keep_next=False, outline=None, numbered=None, rule=False, left=None, hanging=None, right=None):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    st.quick_style = True
    font = st.font
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if small_caps is not None:
        font.small_caps = small_caps
    if underline is not None:
        font.underline = underline
    if color:
        font.color.rgb = RGBColor.from_string(color)
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if keep_next:
        pf.keep_with_next = True
    ppr = st.element.get_or_add_pPr()
    if numbered is not None:
        numbering_ref(ppr, OUTLINE[1], numbered)
    if rule:
        bottom_rule(ppr)
    if left is not None:
        indent(ppr, left, hanging, right)
    if outline is not None:
        ppr_insert(ppr, _el("w:outlineLvl", **{"w:val": outline}))
    arial(st.element.get_or_add_rPr())
    return st


def build_styles(doc):
    part = doc.styles.element
    defaults = part.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            arial(rpr)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.0
    arial(normal.element.get_or_add_rPr())

    LEFT, CENTER = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
    style(doc, "AKLA Title", size=TITLE_PT, bold=True, small_caps=True, color=GOLD, align=CENTER, before=0, after=0)
    style(doc, "AKLA Front", align=LEFT, after=4)
    style(doc, "AKLA Heading 1", size=HEADING_PT, bold=True, small_caps=True, align=LEFT, before=12, after=10,
          keep_next=True, outline=0, numbered=0, rule=True)
    style(doc, "AKLA Heading 2", bold=True, small_caps=True, underline=True, align=LEFT, before=4, after=10,
          keep_next=True, outline=1, numbered=1)
    style(doc, "AKLA Body 1", numbered=1)
    style(doc, "AKLA Body 2", numbered=2)
    style(doc, "AKLA Divider", bold=True, small_caps=True, align=CENTER, before=12, after=10, keep_next=True, rule=True)
    style(doc, "AKLA List", after=6)
    style(doc, "AKLA Quote", italic=True, after=10)
    style(doc, "AKLA Table Text", align=LEFT, before=0, after=0)
    style(doc, "AKLA Table Header", bold=True, small_caps=True, color=GOLD, align=CENTER, before=0, after=0)
    style(doc, "AKLA Header", size=SMALL_PT, small_caps=True, align=LEFT, before=0, after=0)
    style(doc, "AKLA Footer", size=SMALL_PT, align=CENTER, before=0, after=0)


# --------------------------------------------------------------------------
# page furniture
# --------------------------------------------------------------------------

def build_page(doc, args):
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_W, PAGE_H
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, MARGIN)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    # The Guide: the first page is clean - no reference strip, no page
    # number. Every page after carries both.
    section.different_first_page_header_footer = True
    for part in (section.first_page_header, section.first_page_footer):
        part.is_linked_to_previous = False
        part.paragraphs[0].style = doc.styles["AKLA Footer"]

    footer = section.footer.paragraphs[0]
    footer.style = doc.styles["AKLA Footer"]
    for text, field in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        run = footer.add_run(text)
        run.font.size = Pt(SMALL_PT)
        add_field(footer, field)

    build_header(doc, section, args.doc_title, args.doc_status, args.doc_notice, args.doc_date, args.logo)


def build_header(doc, section, title, status, notice, date, logo):
    """The reference strip on the left, the AK emblem on the right."""
    lines = [(line.strip(), True, "000000") for line in (title or "").split("|") if line.strip()]
    if status:
        lines.append((status, False, "000000"))
    if notice:
        lines.append((notice, False, RED))
    if date:
        lines.append((date, False, "000000"))
    header = section.header
    original = header.paragraphs[0]
    if not lines and not logo:
        return
    table = header.add_table(rows=1, cols=2, width=Inches(6.27))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_borders(table, "nil")
    cell_margins(table, 0, 0, 0, 0)
    left, right = table.rows[0].cells
    left.width, right.width = Inches(5.2), Inches(1.07)
    for i, (text, bold, color) in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        p.style = doc.styles["AKLA Header"]
        run = p.add_run(text)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    rp = right.paragraphs[0]
    rp.style = doc.styles["AKLA Header"]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if logo and os.path.exists(logo):
        rp.add_run().add_picture(logo, height=Inches(0.55))
    elif logo:
        print(f"warning: logo not found at {logo}", file=sys.stderr)
    # The header part must end in a paragraph; keep it as small as possible.
    original._element.getparent().remove(original._element)
    tail = header.add_paragraph()
    tail.style = doc.styles["AKLA Header"]
    tail.paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------------------
# comments
# --------------------------------------------------------------------------

# The closing "]]" is the last pair in a run of brackets, so a comment that
# ends on a placeholder - "... PKR [●]" - keeps its own bracket.
COMMENT = re.compile(r"\[\[\s*AKLA(?:\s+Comments?)?\s*:\s*(.+?)\s*\]\](?!\])", re.IGNORECASE)
FOOTNOTE_DEF = re.compile(r"^\[\^([^\]]+)\]:\s*(.*)$")
FOOTNOTE_REF = re.compile(r"\[\^([^\]]+)\](?!:)")
COMMENT_AUTHOR = "AKLA Comments"


def plain_note(text):
    text = re.sub(r"^\*\*\s*AKLA\s+Comments?\s*\d*\s*[.:]?\s*\*\*\s*", "", text.strip(), flags=re.IGNORECASE)
    return re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*", lambda m: m.group(1) or m.group(2), text).strip()


def fold_footnotes(markdown):
    """Markdown footnotes become comments where they are cited.

    Models write the Firm's remarks as footnotes; in a Word document they
    belong in the margin as comments, not in the text or at the end. A
    trailing "AKLA Comments" heading left with nothing under it is dropped.
    """
    lines = markdown.replace("\r\n", "\n").split("\n")
    notes, kept = {}, []
    i = 0
    while i < len(lines):
        m = FOOTNOTE_DEF.match(lines[i].strip())
        if not m:
            kept.append(lines[i])
            i += 1
            continue
        body = [m.group(2)]
        i += 1
        while i < len(lines) and lines[i].startswith(("    ", "\t")) and lines[i].strip():
            body.append(lines[i].strip())
            i += 1
        notes[m.group(1)] = plain_note(" ".join(body))
    if not notes:
        return markdown
    text = FOOTNOTE_REF.sub(lambda m: f"[[AKLA Comment: {notes[m.group(1)]}]]" if m.group(1) in notes else "", "\n".join(kept))
    return re.sub(r"(?im)^#{1,4}\s*AKLA\s+Comments?\s*$\s*(?=^#|\Z)", "", text)


LOG_HEADING = re.compile(r"^#{1,4}\s*(?:\d+[.)]?\s+)?(comments?\s+log|akla\s+comments?|open\s+items(?:\s+log)?|drafting\s+comments)\s*$", re.IGNORECASE)
LOG_ENTRY = re.compile(r"^\s*(?:[-*•]\s*)?(?:\*\*)?\[(C\d+)\](?:\*\*)?\s*[—–:.\-]*\s*(.+)$")
ANCHOR = re.compile(r"`?\[(C\d+)\]`?")


def fold_comments_log(markdown):
    """[C1] anchors and a "Comments Log" section become margin comments.

    Asked for remarks without being told how, models mark each point [C1] in
    the text and list the explanations at the end. In a Word document the
    explanation belongs on the text it concerns: the first mention of each
    anchor takes the comment, later mentions ("cross-referenced") lose the
    marker, and the log section goes.
    """
    lines = markdown.split("\n")
    start = next((i for i, l in enumerate(lines) if LOG_HEADING.match(l.strip())), None)
    if start is None:
        return markdown
    end = next((i for i in range(start + 1, len(lines)) if re.match(r"^#{1,4}\s", lines[i].strip())), len(lines))
    notes = {}
    for line in lines[start + 1:end]:
        m = LOG_ENTRY.match(line)
        if m:
            notes[m.group(1)] = plain_note(m.group(2))
    if not notes:
        return markdown
    body = "\n".join(lines[:start] + lines[end:])
    used = set()

    def swap(m):
        key = m.group(1)
        if key not in notes:
            return m.group(0)
        if key in used:
            return ""
        used.add(key)
        return f"[[AKLA Comment: {notes[key]}]]"

    body = ANCHOR.sub(swap, body)
    body = re.sub(r"\s*,?\s*\((?:cross[- ]referenced|see above)\)", "", body, flags=re.IGNORECASE)
    body = re.sub(r"(\]\])\s*,\s*(?=\[\[|$)", r"\1 ", body, flags=re.MULTILINE)
    return body.rstrip() + "\n"


TYPED_NUMBER = re.compile(r"^(?:\d+(?:\.\d+)+\.?|[A-Z](?:\.\d+)+\.?)\s+(?=\S)")
TOC_HEADING = re.compile(r"^(?:table\s+of\s+)?contents$", re.IGNORECASE)


def take_comments(text):
    """The text without its comment markers, and the comments it carried."""
    notes = [m.group(1).strip() for m in COMMENT.finditer(text)]
    return COMMENT.sub("", text).rstrip(), notes


def attach_comments(doc, paragraph, notes):
    if not notes or paragraph is None:
        return
    runs = [r for r in paragraph.runs if r.text]
    if not runs:
        return
    for note in notes:
        doc.add_comment(runs, text=note, author=COMMENT_AUTHOR, initials="AKLA")


# --------------------------------------------------------------------------
# text
# --------------------------------------------------------------------------

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|_[^_\s][^_]*?_)")


def add_runs(paragraph, text):
    """Runs for a line: **bold**, *italic*, and the firm's name in italic navy."""
    for token in filter(None, INLINE.split(text)):
        bold = token.startswith("**") and token.endswith("**") and len(token) > 4
        italic = not bold and len(token) > 2 and token[0] in "*_" and token[-1] == token[0]
        inner = token[2:-2] if bold else token[1:-1] if italic else token
        for piece in filter(None, FIRM_NAME.split(inner)):
            run = paragraph.add_run(piece)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if FIRM_NAME.fullmatch(piece):
                # The Guide: the firm's name in running text is italic, dark blue.
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(NAVY)


def banner(doc, text):
    lines = text if isinstance(text, list) else [text]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table, "nil")
    table_width(table, TEXT_TWIPS)
    cell_margins(table, 288, 288, 144, 144)
    cell = table.rows[0].cells[0]
    cell_fill(cell, NAVY)
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.style = doc.styles["AKLA Title"]
        add_runs(p, line)
    doc.add_paragraph(style="AKLA Front").paragraph_format.space_after = Pt(6)


def box(doc, lines, indent_twips):
    table = doc.add_table(rows=1, cols=1)
    table_borders(table, "nil")
    table_width(table, TEXT_TWIPS - indent_twips, indent_twips)
    cell_margins(table, 144, 144, 144, 144)
    cell = table.rows[0].cells[0]
    cell_fill(cell, SHADE)
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = doc.styles["AKLA Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_start = len(p.runs)
        add_runs(p, line)
        for run in p.runs[run_start:]:
            run.bold = True
            run.font.small_caps = True
    doc.add_paragraph(style="AKLA Front").paragraph_format.space_after = Pt(4)


def data_table(doc, rows, indent_twips):
    header, body = rows[0], rows[1:]
    cols = len(header)
    table = doc.add_table(rows=len(rows), cols=cols)
    table_borders(table)
    width = TEXT_TWIPS - indent_twips
    table_width(table, width, indent_twips)
    cell_margins(table, 144, 144)
    grid = table._tbl.find(qn("w:tblGrid"))
    for col in grid.findall(qn("w:gridCol")):
        col.set(qn("w:w"), str(width // cols))
    trpr = table.rows[0]._tr.get_or_add_trPr()
    _insert(trpr, _el("w:tblHeader"), TRPR_ORDER)
    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.width = width * 635 // cols
        cell_fill(cell, NAVY)
        p = cell.paragraphs[0]
        p.style = doc.styles["AKLA Table Header"]
        add_runs(p, take_comments(text)[0])
    for i, row in enumerate(body, start=1):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.width = width * 635 // cols
            p = cell.paragraphs[0]
            p.style = doc.styles["AKLA Table Text"]
            text, notes = take_comments(row[j] if j < len(row) else "")
            add_runs(p, text)
            attach_comments(doc, p, notes)
    doc.add_paragraph(style="AKLA Front").paragraph_format.space_after = Pt(4)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


DIVIDER = re.compile(r"^\|[\s:|-]+\|$")
ITEM = re.compile(r"^(\s*)(?:[-*•]|\d{1,2}[.)]|\([a-z]{1,4}\))\s+(.*)$")


def render(doc, markdown, numbering):
    lists = ListNumbers(numbering)
    title_done = False
    in_sections = False
    body_level = 1        # the outline level body text takes where it stands
    list_num = None       # the w:num of the list in progress, if any
    last = None
    lines = fold_comments_log(fold_footnotes(markdown.replace("\r\n", "\n"))).split("\n")
    lines = [re.sub(r"`([^`]*)`", r"\1", l) for l in lines]   # code marks are not document formatting
    i = 0
    while i < len(lines):
        raw = lines[i]
        line, notes = take_comments(raw.strip())
        i += 1
        if not line:
            attach_comments(doc, last, notes)
            continue

        text_indent = 0   # nothing is indented; tables and boxes run from the margin

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", line):
            # A Markdown rule is a visual separator, not a paragraph.
            attach_comments(doc, last, notes)
            list_num = None
            continue

        if line.startswith("|") and i < len(lines) and DIVIDER.match(lines[i].strip()):
            rows = [split_row(line)]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            data_table(doc, rows, text_indent if in_sections else 0)
            list_num = None
            continue

        if line.startswith(":::"):
            group = [line[3:].strip()]
            while i < len(lines) and lines[i].strip().startswith(":::"):
                group.append(lines[i].strip()[3:].strip())
                i += 1
            box(doc, group, text_indent if in_sections else 0)
            list_num = None
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            depth, text = len(heading.group(1)), heading.group(2).strip()
            text = re.sub(r"^(\d+(\.\d+)*\.?|[A-Z](\.\d+)*\.|[A-Z]\.\d+(\.\d+)*|[IVX]+\.)\s+", "", text)  # typed numbers
            list_num = None
            if depth == 1:
                # "## " lines straight after the title, with nothing between
                # them, are the rest of the title set on separate lines.
                parts = [text]
                j = i
                while j < len(lines) and not lines[j].strip():
                    j += 1
                run = []
                while j < len(lines) and re.match(r"^##\s+", lines[j].strip()):
                    run.append(lines[j].strip()[3:].strip())
                    j += 1
                if len(run) >= 2:
                    parts += run
                    i = j
                banner(doc, parts)
                title_done = True
                if notes:
                    print("warning: a comment on the title was dropped", file=sys.stderr)
                continue
            if TOC_HEADING.match(text):
                # A contents list typed by hand goes stale the moment a
                # section moves; Word's own table of contents does not.
                p = doc.add_paragraph(style="AKLA Divider")
                add_runs(p, "Table Of Contents")
                toc = doc.add_paragraph(style="AKLA Front")
                field = _el("w:fldSimple", **{"w:instr": 'TOC \\o "1-2" \\h \\z \\u'})
                placeholder = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = "Right-click and choose Update Field to show the table of contents."
                placeholder.append(t)
                field.append(placeholder)
                toc._p.append(field)
                while i < len(lines) and not re.match(r"^#{1,4}\s", lines[i].strip()):
                    i += 1
                last = p
                continue
            if depth == 4:
                p = doc.add_paragraph(style="AKLA Divider")
            else:
                in_sections = True
                p = doc.add_paragraph(style="AKLA Heading 1" if depth == 2 else "AKLA Heading 2")
                body_level = 1 if depth == 2 else 2
            add_runs(p, text)
            attach_comments(doc, p, notes)
            last = p
            continue

        quote = re.match(r"^>\s?(.*)$", line)
        if quote:
            p = doc.add_paragraph(style="AKLA Quote")
            # Quoted text lines up with the text around it; italics mark it.
            indent(p._p.get_or_add_pPr(), TEXT_AT if in_sections else 0)
            add_runs(p, quote.group(1))
            attach_comments(doc, p, notes)
            last = p
            continue

        item = ITEM.match(take_comments(raw.rstrip())[0]) if in_sections else None
        if item:
            depth = min(2, len(item.group(1).replace("\t", "    ")) // 2)
            if list_num is None:
                list_num = lists.new(LIST_UNDER_BODY1 if body_level == 1 else LIST_UNDER_BODY2)
            p = doc.add_paragraph(style="AKLA List")
            numbering_ref(p._p.get_or_add_pPr(), list_num, depth)
            add_runs(p, item.group(2).strip())
            attach_comments(doc, p, notes)
            last = p
            continue

        list_num = None
        if not in_sections:
            p = doc.add_paragraph(style="AKLA Front")
        else:
            p = doc.add_paragraph(style="AKLA Body 1" if body_level == 1 else "AKLA Body 2")
            line = TYPED_NUMBER.sub("", line)   # Word numbers the clause itself
        add_runs(p, line)
        attach_comments(doc, p, notes)
        last = p

    if not title_done:
        print("warning: no '# Title' line found", file=sys.stderr)


def main():
    default_logo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "akla-logo.png")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="Markdown file to render, or - for stdin")
    ap.add_argument("-o", "--out", required=True, help="output .docx path")
    ap.add_argument("--doc-title", default="", help="header reference strip, bold; use | to split lines")
    ap.add_argument("--doc-status", default="", help="header version line, e.g. 'First Circulation Version'")
    ap.add_argument("--doc-notice", default="Privileged And Confidential", help="header confidentiality line, printed in red")
    ap.add_argument("--doc-date", default="", help="header date line")
    ap.add_argument("--logo", default=default_logo, help="header emblem (default: the bundled AK mark)")
    ap.add_argument("--no-logo", action="store_true")
    ap.add_argument("--heading-pt", type=int, default=HEADING_PT, help=argparse.SUPPRESS)
    args = ap.parse_args()
    if args.no_logo:
        args.logo = ""

    if args.source == "-":
        markdown = sys.stdin.read()
    else:
        with open(args.source, encoding="utf-8") as fh:
            markdown = fh.read()

    doc = Document()
    numbering = build_numbering(doc)
    build_styles(doc)
    build_page(doc, args)
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)   # the template's empty first paragraph
    render(doc, markdown, numbering)
    if doc.element.body.find(".//" + qn("w:fldSimple")) is not None:
        settings = doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            settings.append(_el("w:updateFields", **{"w:val": "true"}))
    doc.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
