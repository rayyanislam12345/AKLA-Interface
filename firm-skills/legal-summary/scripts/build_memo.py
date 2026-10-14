"""Build an AKLA "Notes On ..." memo from Markdown, in house format.

Renders with this skill's build_docx.py (navy/gold headings, 1. / 1.1. /
1.1.1. / (a) Word numbering, running header with the AK emblem), then applies
the house spacing rule: space before/after is 0 on every paragraph and style,
and each gap the builder made with spacing becomes an empty paragraph.

Usage:
  python3 tools/build_memo.py work/share_note.md "output/<name>.docx" \
      --title "Notes On ..." --status "First Draft" --date "October 06, 2026"

Adds widen_quotes() for long excerpts and
{{citation|excerpt}} markup: the citation prints in the text and the verbatim excerpt
is attached to it as a Word comment.
"""
import argparse
import copy
import re
import tempfile
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.text.run import Run
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches

SKILL_BUILDER = Path(__file__).resolve().parent / "build_docx.py"


def effective(p, attr):
    v = getattr(p.paragraph_format, attr)
    st = p.style
    while v is None and st is not None:
        v = getattr(st.paragraph_format, attr)
        st = st.base_style
    return v or 0


def is_empty(el):
    return el is not None and el.tag == qn("w:p") and not "".join(el.itertext()).strip() \
        and el.find(".//" + qn("w:numPr")) is None


def blank():
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    ppr.append(sp)
    p.append(ppr)
    return p


def _is_item(el):
    """A lettered list item: numPr at level 3 (the builder's (a) level)."""
    np_ = el.find(".//" + qn("w:numPr"))
    if np_ is None:
        return False
    il = np_.find(qn("w:ilvl"))
    return il is not None and il.get(qn("w:val")) == "3"


def zero_spacing(doc, tight_lists=False):
    body = doc.element.body
    paras = [p for p in doc.paragraphs]          # body-level paragraphs only
    gaps = [(p._p, effective(p, "space_before"), effective(p, "space_after")) for p in paras]
    for el, before, after in gaps:
        if tight_lists and _is_item(el) and el.getnext() is not None and _is_item(el.getnext()):
            after = 0                             # no blank line between consecutive bullets
        if before and not is_empty(el.getprevious()) and el.getprevious() is not None:
            el.addprevious(blank())
        if after and not is_empty(el.getnext()) and el.getnext() is not None \
                and el.getnext().tag != qn("w:sectPr"):
            el.addnext(blank())
    # zero spacing everywhere: paragraphs (body and tables) and styles
    for p in body.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        sp = ppr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            ppr.append(sp)
        sp.set(qn("w:before"), "0")
        sp.set(qn("w:after"), "0")
    for st in doc.styles:
        if hasattr(st, "paragraph_format"):
            st.paragraph_format.space_before = 0
            st.paragraph_format.space_after = 0


CITE = re.compile(r"\{\{([^|{}]+)\|([^|{}]+)(?:\|([^|{}]+))?\}\}")


def mark_citations(markdown):
    """{{citation|excerpt}} -> citation wrapped in numbered markers; excerpts returned."""
    notes = []

    def sub(m):
        notes.append(((m.group(3) or m.group(1)).strip(), m.group(2).strip()))
        n = len(notes) - 1
        return f"\u27e6{n}\u27e7{m.group(1).strip()}\u27e6/{n}\u27e7"
    return CITE.sub(sub, markdown), notes


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def attach_comments(doc, notes, default_act=None, author="AKLA", initials="AK"):
    """Strip the markers and hang each excerpt on its citation as a Word comment."""
    marker = re.compile(r"(\u27e6/?\d+\u27e7)")
    count = 0
    for p in list(all_paragraphs(doc)):
        if "\u27e6" not in p.text:
            continue
        spans, open_n, start = {}, None, None
        last = None
        for run in list(p.runs):
            if "\u27e6" not in run.text:
                if open_n is not None and start is None:
                    start = run
                last = run
                continue
            parts = [x for x in marker.split(run.text) if x != ""]
            anchor = run._r
            made = []
            for part in parts:
                m = re.fullmatch(r"\u27e6(/?)(\d+)\u27e7", part)
                if m:
                    if m.group(1):
                        spans[int(m.group(2))] = (start, last)
                        open_n, start = None, None
                    else:
                        open_n, start = int(m.group(2)), None
                    continue
                r = copy.deepcopy(run._r)
                anchor.addnext(r)
                anchor = r
                nr = Run(r, p)
                nr.text = part
                if open_n is not None and start is None:
                    start = nr
                last = nr
            run._r.getparent().remove(run._r)
        for n, (a, b) in spans.items():
            cit, exc = notes[n]
            if default_act and re.match(r"s\.\d", cit):
                cit = f"{default_act}, {cit}"
            doc.add_comment([a, b], text=f'{cit}: "{exc}"', author=author, initials=initials)
            count += 1
    return count


def widen_quotes(doc):
    """Excerpts render at 1.75" left indent in the skill; this note quotes long
    statutory text, so pull them in to 1.0" / 0.5"."""
    for p in doc.paragraphs:
        if p.paragraph_format.left_indent == Inches(1.75):
            p.paragraph_format.left_indent = Inches(1.0)
            p.paragraph_format.right_indent = Inches(0.5)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("out")
    ap.add_argument("--title", required=True)
    ap.add_argument("--status")
    ap.add_argument("--date", required=True)
    ap.add_argument("--default-act", help='prefix for bare "s.N" citations in comments')
    ap.add_argument("--tight-lists", action="store_true", help="no blank line between consecutive bullets (one-pagers)")
    a = ap.parse_args()
    marked, notes = mark_citations(Path(a.source).read_text(encoding="utf-8"))
    tmp = tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8")
    tmp.write(marked)
    tmp.close()
    cmd = [sys.executable, str(SKILL_BUILDER), tmp.name, "-o", a.out, "--doc-title", a.title, "--doc-date", a.date]
    if a.status:
        cmd += ["--doc-status", a.status]
    subprocess.run(cmd, check=True)
    doc = Document(a.out)
    widen_quotes(doc)
    zero_spacing(doc, a.tight_lists)
    n = attach_comments(doc, notes, a.default_act)
    doc.save(a.out)
    Path(tmp.name).unlink()
    left = sum("\u27e6" in p.text for p in all_paragraphs(Document(a.out)))
    if n != len(notes) or left:
        sys.exit(f"comment mismatch: {n} attached of {len(notes)}, {left} paragraphs with stray markers")
    print(f"built {a.out} (spacing zeroed, gaps as blank paragraphs, {n} excerpt comments)")


if __name__ == "__main__":
    main()
