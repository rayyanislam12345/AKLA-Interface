#!/usr/bin/env python3
"""Audit a DD volume against the AKLA gold standard, and prove a reformat lost no wording.

    python3 check_volume.py "Volume III.docx"                       # format audit
    python3 check_volume.py new.docx --against "original draft.docx" # + content-loss check

The second form is the one that matters. Reformatting a legal document is only safe if you
can show the words came through unchanged, so run it on every rebuild before handing the
file back. Exit code is 1 if anything is missing or a hard rule fails.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

NAVY, GOLD, GREY, BLUE, RED, GREEN = "002060", "FFC000", "D9D9D9", "0070C0", "FF0000", "00B050"
QUOTES = {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-", " ": " "}
RUNNING = re.compile(r"^(legal due diligence report|volume [ivx]+ \(|first draft|second draft|final|"
                     r"strictly confidential|page \d+ of \d+)", re.I)
MARKER = re.compile(r"(?:(?<=\s)|^)\((?:[a-z]{1,4}|[ivx]{1,6}|\d{1,2})\)(?=\s|$)")


def norm(s: str) -> str:
    for a, b in QUOTES.items():
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s.lower()).strip()


def key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", MARKER.sub(" ", " " + norm(s))).strip()


def sentences(text: str, min_words=4) -> list[str]:
    out = []
    for c in re.split(r"(?<=[.;:?!])\s+(?=[A-Z(‘\"']|\d+\.)", text):
        c = re.sub(r"\s+", " ", c).strip()
        if len(c.split()) >= min_words:
            out.append(c)
    return out


def walk(path: Path):
    """('para'|'table', obj) in document order, recursing into table cells."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))

    def inner(el, parent):
        for child in el.iterchildren():
            tag = child.tag.split("}")[1]
            if tag == "p":
                yield "para", Paragraph(child, parent)
            elif tag == "tbl":
                t = Table(child, parent)
                yield "table", t
                for row in t.rows:
                    for cell in row.cells:
                        yield from inner(cell._tc, cell)
    yield from inner(d.element.body, d)


def units(path: Path) -> list[str]:
    out = []
    for kind, obj in walk(path):
        if kind != "para":
            continue
        t = re.sub(r"\s+", " ", obj.text).strip()
        if not t or RUNNING.match(t) or t.startswith("[***"):
            continue
        out.extend(sentences(t))
    return out


def full_text(path: Path) -> str:
    return key(" ".join(re.sub(r"\s+", " ", p.text) for k, p in walk(path) if k == "para"))


def coverage(unit_key: str, blob: str, min_words=4):
    """Longest in-order fragments of the unit that appear in blob; returns (coverage, biggest gap)."""
    w = unit_key.split()
    covered, gaps, gap, i = 0, [], 0, 0
    while i < len(w):
        j = len(w)
        while j > i:
            if (j - i) >= min_words and " ".join(w[i:j]) in blob:
                covered += j - i
                break
            j -= 1
        if j == i:
            gap += 1
            i += 1
        else:
            if gap:
                gaps.append(gap)
            gap = 0
            i = j
    if gap:
        gaps.append(gap)
    return covered / max(1, len(w)), max(gaps or [0])


# ---------------------------------------------------------------- format audit
def audit(path: Path) -> list[tuple[str, str]]:
    """Returns [(level, message)] where level is FAIL or WARN."""
    import docx
    from docx.shared import RGBColor
    d = docx.Document(str(path))
    out = []
    sec = d.sections[0]
    w, h = sec.page_width.inches, sec.page_height.inches
    # A4 portrait is the house format from September 14, 2026. Landscape remains acceptable only for a wide
    # schedule that cannot be restructured, or for a re-issue of a volume already circulated in landscape.
    if w > h:
        out.append(("WARN", f"page is landscape ({w:.2f}\" x {h:.2f}\") - DD volumes are now A4 portrait 8.27\" x 11.69\"; "
                            "landscape only for a wide schedule or a re-issue of a volume circulated in landscape"))
    elif abs(w - 8.27) > 0.15 or abs(h - 11.69) > 0.15:
        out.append(("WARN", f"page is {w:.2f}\" x {h:.2f}\", expected A4 portrait 8.27\" x 11.69\""))
    for name, got, want in (("left", sec.left_margin.inches, 0.8), ("right", sec.right_margin.inches, 0.8),
                            ("top", sec.top_margin.inches, 1.35), ("bottom", sec.bottom_margin.inches, 0.75)):
        if abs(got - want) > 0.12:
            out.append(("WARN", f"{name} margin {got:.2f}\", expected {want}\""))

    base = d.styles["Normal"].font.name
    if base and base != "Arial":
        out.append(("FAIL", f"document default font is {base}, expected Arial"))
    fonts = {}
    for k, p in walk(path):
        if k != "para":
            continue
        for r in p.runs:
            if r.text.strip() and r.font.name:
                fonts[r.font.name] = fonts.get(r.font.name, 0) + len(r.text)
    if fonts:
        main = max(fonts, key=fonts.get)
        if main != "Arial":
            out.append(("FAIL", f"body font is {main}, expected Arial"))
        other = {f: n for f, n in fonts.items() if f != "Arial" and n > 200}
        if other:
            out.append(("WARN", f"non-Arial runs present: {other}"))

    hdr_text = " ".join(p.text for p in sec.header.paragraphs)
    for t in sec.header.tables:
        for row in t.rows:
            for c in row.cells:
                hdr_text += " " + " ".join(p.text for p in c.paragraphs)
    if not hdr_text.strip():
        out.append(("FAIL", "no running header - needs the five-line reference strip and the AK emblem"))
    else:
        if "CONFIDENTIAL" not in hdr_text.upper():
            out.append(("WARN", "header has no confidentiality line"))
        red = any(r.font.color and r.font.color.rgb and str(r.font.color.rgb) == RED
                  for t in sec.header.tables for row in t.rows for c in row.cells
                  for p in c.paragraphs for r in p.runs)
        if not red:
            out.append(("WARN", f"header confidentiality line is not red ({RED})"))
    if not any("PAGE" in (p.text or "") or p.runs for p in sec.footer.paragraphs):
        out.append(("WARN", "no footer - expected 'Page X of Y' centered, Arial 8pt"))

    text_all = " ".join(p.text for k, p in walk(path) if k == "para")
    n_comment = len(re.findall(r"AKLA Comment\s*:", text_all))
    n_action = len(re.findall(r"AKLA Action Item\s*:", text_all))
    if n_comment == 0 and n_action == 0:
        out.append(("WARN", "no AKLA Comments or Action Items found - unusual for a DD volume"))
    wrong = []
    for k, p in walk(path):
        if k != "para":
            continue
        t = p.text.strip()
        m = re.match(r"^AKLA (Comment|Action Item|Response)\s*:", t)
        if not m:
            continue
        want = {"Comment": BLUE, "Action Item": RED, "Response": GREEN}[m.group(1)]
        got = next((str(r.font.color.rgb) for r in p.runs if r.font.color and r.font.color.rgb), None)
        if got != want:
            wrong.append(f"{m.group(1)} is {got or 'default/black'}, expected {want}")
    if wrong:
        seen = sorted(set(wrong))
        out.append(("FAIL", f"{len(wrong)} AKLA paragraph(s) mis-coloured: {'; '.join(seen[:3])}"))

    heads = [p.text.strip() for k, p in walk(path) if k == "para" and p.style.name.startswith("Heading") and p.text.strip()]
    if not heads:
        out.append(("WARN", "no Heading-styled paragraphs - section headings should use Heading 1/2 so the TOC builds"))
    body = "\n".join(p.text for k, p in walk(path) if k == "para")
    for label, pat in (("Important Note", r"IMPORTANT NOTE"), ("Glossary", r"GLOSSARY"),
                       ("Annexure A", r"ANNEXURE A")):
        if not re.search(pat, body, re.I):
            out.append(("WARN", f"no {label} section found"))
    if not re.search(r"TABLE OF CONTENTS", body, re.I):
        out.append(("WARN", "no table of contents"))

    # zero spacing before/after, gaps as blank paragraphs
    spaced = 0
    for p in d.element.body.iter(qn_w("p")):
        sp = p.find(".//" + qn_w("spacing"))
        if sp is not None and (sp.get(qn_w("before"), "0") not in ("0",) or sp.get(qn_w("after"), "0") not in ("0",)):
            spaced += 1
    if spaced:
        out.append(("WARN", f"{spaced} paragraph(s) carry space before/after - house rule is 0 and 0, gaps as blank lines"))
    typed = [p.text.strip()[:40] for k, p in walk(path) if k == "para" and re.match(r"^\((?:[a-z]|[ivxl]+)\)\t", p.text)]
    if typed:
        out.append(("WARN", f"{len(typed)} list item(s) with typed labels such as '{typed[0]}' - letters should be Word's list numbering"))
    page_w = w - sec.left_margin.inches - sec.right_margin.inches
    for t in d.tables:
        if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip().upper() == "DEFINED TERM" and len(t.rows) > 1:
            tw = t.rows[1].cells[0].width
            if tw and page_w and tw.inches / page_w > 0.35:
                out.append(("WARN", f"glossary term column is {tw.inches / page_w:.0%} of the text width - use about 29% / 71%"))
            break
    if re.search(r"\u2014\s*$", body, re.M) and re.search(r"TABLE OF CONTENTS", body, re.I):
        out.append(("WARN", "contents page numbers not yet refreshed ('—') - update fields in Word before export"))

    for t in d.tables:
        if not t.rows:
            continue
        hdr = t.rows[0]
        fill = hdr.cells[0]._tc.find(".//" + qn_w("shd"))
        got = fill.get(qn_w("fill")) if fill is not None else None
        if got and got.upper() not in (NAVY, GREY):
            out.append(("WARN", f"table header shaded {got}, expected navy {NAVY} or grey {GREY}"))
            break
    return out


def qn_w(tag: str) -> str:
    return "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag


# ---------------------------------------------------------------- content check
def compare(original: Path, rebuilt: Path):
    o_units = units(original)
    blob = full_text(rebuilt)
    r_keys = {key(u) for u in units(rebuilt)}
    missing, reflowed, exact = [], 0, 0
    for u in o_units:
        k = key(u)
        if not k:
            continue
        if k in r_keys or k in blob:
            exact += 1
            continue
        cov, gap = coverage(k, blob)
        absent = [w for w in dict.fromkeys(k.split()) if w not in blob]
        if not absent and gap <= 1 and cov >= 0.90:
            reflowed += 1
        else:
            missing.append((u, round(cov, 2), absent[:8]))
    ow = set(key(" ".join(o_units)).split())
    words = len(ow & set(blob.split())) / max(1, len(ow))
    return {"total": len([u for u in o_units if key(u)]), "exact": exact, "reflowed": reflowed,
            "missing": missing, "word_coverage": words}


def main():
    ap = argparse.ArgumentParser(description="Audit a DD volume's formatting, and verify a reformat lost no wording.")
    ap.add_argument("file")
    ap.add_argument("--against", help="the original draft, to prove nothing was dropped")
    a = ap.parse_args()
    path = Path(a.file)
    if not path.exists():
        raise SystemExit(f"not found: {path}")

    print(f"FORMAT AUDIT: {path.name}")
    issues = audit(path)
    if not issues:
        print("  matches the gold standard on every checked point")
    for lvl, msg in issues:
        print(f"  {lvl}: {msg}")

    bad = any(l == "FAIL" for l, _ in issues)
    if a.against:
        orig = Path(a.against)
        print(f"\nCONTENT CHECK against {orig.name}")
        r = compare(orig, path)
        print(f"  {r['exact'] + r['reflowed']}/{r['total']} original text units present "
              f"({r['exact']} verbatim, {r['reflowed']} re-flowed) | word coverage {r['word_coverage']:.2%}")
        if r["missing"]:
            bad = True
            print(f"  MISSING {len(r['missing'])} unit(s) - DO NOT send this file:")
            for u, cov, absent in r["missing"][:10]:
                print(f"    cov={cov} absent={absent}\n      {u[:150]}")
        else:
            print("  nothing lost")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
