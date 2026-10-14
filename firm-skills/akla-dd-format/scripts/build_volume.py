#!/usr/bin/env python3
"""Render a DD volume structure (JSON) into a Word file in AKLA gold-standard format.

    python3 build_volume.py volume.json -o "SP3A - DD Volume III [AKLA][First Draft][August 27, 2026].docx"

The JSON is what read_volume.py emits. Every colour, font and numbering value here was measured off the
firm's circulated volumes; see references/dd-volume-format.md.

Pages are A4 PORTRAIT by default (the firm's format from September 14, 2026). Margins, header/footer
distances, fonts, colours and table point sizes are unchanged from the circulated landscape volumes; only the
page turns and widths are fitted to the narrower text block. Pass --landscape (or "orientation": "landscape"
in the JSON) only for a genuinely wide schedule that cannot be restructured, or to re-issue a volume already
circulated in landscape.

Only python-docx is required.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY, GOLD, GREY = "002060", "FFC000", "D9D9D9"
BLUE, RED, BLACK, GREEN = "0070C0", "FF0000", "000000", "00B050"
FONT = "Arial"
ALPHA = "abcdefghijklmnopqrstuvwxyz"

# Optional sections (development package, September 15, 2026). A volume JSON without these keys renders as before.
METHODOLOGY_TITLE = "SCOPE, TRANSACTION OVERVIEW AND METHODOLOGY"
SUMMARY_TITLE = "EXECUTIVE SUMMARY"
ACTIONS_TITLE = "TRANSACTION ACTIONS"
COVERAGE_TITLE = "ANNEXURE B – REVIEW COVERAGE"
ASSETS = Path(__file__).resolve().parent.parent / "assets"


def _roman(n: int) -> str:
    out = ""
    for v, s in ((10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")):
        while n >= v:
            out += s
            n -= v
    return out


# ---------------------------------------------------------------- low-level XML
def _shade(cell, fill: str):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100):
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


def _borders(table, sz=4):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    table._tbl.tblPr.append(b)


def _widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)
    for i, gc in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(widths[i] * 1440)))


def _repeat_header(row):
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(e)


def _field(paragraph, instr: str, placeholder="1"):
    r = paragraph.add_run()
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); r._r.append(f)
    r = paragraph.add_run()
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "; r._r.append(t)
    r = paragraph.add_run()
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "separate"); r._r.append(f)
    r = paragraph.add_run(placeholder); r.font.name = FONT; r.font.size = Pt(8)
    r2 = paragraph.add_run()
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r2._r.append(f)


def _bottom_rule(paragraph, sz=6):
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "000000")
    pbdr.append(b)
    paragraph._p.get_or_add_pPr().append(pbdr)


def _style(run, size=11, bold=None, italic=None, color=None, caps=None, underline=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if caps is not None:
        run.font.small_caps = caps
    if underline is not None:
        run.font.underline = underline


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")


def _rich(paragraph, text: str, size=11, color=None, bold=None, italic=None, caps=None):
    if "[***" in (text or ""):                 # the rest-of-page marker is literal text, not markup
        _style(paragraph.add_run(text), size=size, bold=bold, italic=italic, color=color, caps=caps)
        return
    for part in INLINE.split(text or ""):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        _style(paragraph.add_run(part), size=size, bold=b, italic=i, color=color, caps=caps)


# ---------------------------------------------------------------- the volume
def _akla_numbering(doc) -> int:
    """Define (once per document) the firm's multi-level list and return its numId.

    AKLA numbers by Word's own list definition rather than by typing digits, so inserting or deleting a
    paragraph renumbers the rest. Levels render 1. / 1.1. / 1.1.1., text hanging at 0.5in throughout.
    """
    try:
        numbering = doc.part.numbering_part.element
    except (AttributeError, KeyError, NotImplementedError):
        return 0
    found = _named_num(numbering, "AKLA Report Numbering")
    if found:
        return found
    return _new_multilevel(doc, "AKLA Report Numbering")


def _named_num(numbering, name: str) -> int:
    """numId of the list whose abstract definition carries w:name == name (0 if none). Cached in the XML
    itself, so it survives python-docx versions whose Document objects refuse new attributes."""
    for a in numbering.findall(qn("w:abstractNum")):
        n = a.find(qn("w:name"))
        if n is not None and n.get(qn("w:val")) == name:
            aid = a.get(qn("w:abstractNumId"))
            for num in numbering.findall(qn("w:num")):
                ref = num.find(qn("w:abstractNumId"))
                if ref is not None and ref.get(qn("w:val")) == aid:
                    return int(num.get(qn("w:numId")))
    return 0


def _add_abstract(numbering, abstract):
    """OOXML requires every w:abstractNum to precede every w:num."""
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(abstract)
    else:
        numbering.append(abstract)


def _new_multilevel(doc, name: str = "AKLA Numbering") -> int:
    """A fresh 1. / 1.1. / 1.1.1. list (e.g. for Annexure A, which restarts at 1)."""
    numbering = doc.part.numbering_part.element
    aid = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))], default=-1) + 1
    nid = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))], default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(aid))
    mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "multilevel"); abstract.append(mlt)
    nm = OxmlElement("w:name"); nm.set(qn("w:val"), name); abstract.append(nm)
    for i in range(9):
        lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), str(i))
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"),
                         ("w:lvlText", ".".join(f"%{k + 1}" for k in range(i + 1)) + "."),
                         ("w:lvlJc", "left")):
            e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "720")
        ppr.append(ind); lvl.append(ppr); abstract.append(lvl)
    _add_abstract(numbering, abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(aid)); num.append(ref)
    numbering.append(num)
    return nid


def _new_list(doc, style: str = "alpha") -> int:
    """A separate single-level Word list, starting at (a) / (i) / 1., for one lettered list. The letters are
    Word's, never typed."""
    numbering = doc.part.numbering_part.element
    aid = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))], default=-1) + 1
    nid = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))], default=0) + 1
    fmt = {"alpha": "lowerLetter", "roman": "lowerRoman", "number": "decimal"}.get(style, "lowerLetter")
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(aid))
    mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "singleLevel"); abstract.append(mlt)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", "%1." if style == "number" else "(%1)"),
                     ("w:lvlJc", "left"), ("w:suff", "tab")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
    abstract.append(lvl)
    _add_abstract(numbering, abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(aid)); num.append(ref)
    ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
    numbering.append(num)
    return nid


def _bookmark(p, bid: int, name: str):
    b = OxmlElement("w:bookmarkStart"); b.set(qn("w:id"), str(bid)); b.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    p._p.insert(0, b); p._p.append(e)


def _tc(p, text: str):
    """Hidden TC field so the contents lists exactly the firm's landmarks (front matter, level-1 sections,
    annexure) and not every sub-heading."""
    for kind in ("begin", "instr", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve")
            el.text = ' TC "' + text.replace('"', "") + '" \\f C \\l 1 '
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r = p.add_run(); r.font.hidden = True; r._r.append(el)


def _zero_spacing(part_element):
    """House rule: 0 before / 0 after on every paragraph, including inside table cells."""
    for p in part_element.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
        sp = ppr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing"); ppr.append(sp)
        sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")


def _set_num(p, num_id: int, ilvl: int):
    if not num_id:
        return
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    npr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", str(ilvl)), ("w:numId", str(num_id))):
        e = OxmlElement(tag); e.set(qn("w:val"), val); npr.append(e)
    ppr.append(npr)


class Volume:
    def __init__(self, meta: dict):
        self.m = meta
        self.doc = Document()
        self.bookmarks: list[str] = []
        self._page()
        self._styles()
        self._header_footer()

    def _page(self):
        s = self.doc.sections[0]
        if str(self.m.get("orientation", "portrait")).lower() == "landscape":
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Inches(11.69), Inches(8.27)
        else:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin, s.bottom_margin = Inches(1.35), Inches(0.75)
        s.header_distance, s.footer_distance = Inches(0.45), Inches(0.4)
        self.width = (s.page_width - s.left_margin - s.right_margin) / 914400

    def _styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        # House rule: no paragraph spacing; a gap is an explicit empty line.
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        for lvl, name in ((1, "Heading 1"), (2, "Heading 2")):
            h = self.doc.styles[name]
            h.font.name = FONT; h.font.size = Pt(11)
            h.font.bold = lvl == 1
            h.font.small_caps = True; h.font.underline = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before = Pt(0)
            h.paragraph_format.space_after = Pt(0)
            h.paragraph_format.keep_with_next = True
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "false")   # no update prompt on open; refresh fields before export
        self.doc.settings.element.append(upd)

    def _header_footer(self):
        sec = self.doc.sections[0]
        hdr = sec.header
        hdr.is_linked_to_previous = False
        t = hdr.add_table(rows=1, cols=2, width=Inches(self.width))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _widths(t, [self.width - 1.1, 1.1])
        left, right = t.rows[0].cells
        lines = [(self.m["report_title"], True, BLACK),
                 (f"VOLUME {self.m['volume_roman']} ({self.m['volume_title'].upper()})", True, BLACK),
                 (self.m.get("draft_label", "FIRST DRAFT"), False, BLACK),
                 (self.m.get("confidentiality", "STRICTLY CONFIDENTIAL AND FOR DISCUSSION PURPOSES ONLY"), False, RED),
                 (self.m.get("date", ""), False, BLACK)]
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for i, (txt, bold, col) in enumerate([l for l in lines if l[0]]):
            if i:
                p = left.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
            _style(p.add_run(txt), size=6.5, bold=bold, color=col, caps=True)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        logo = ASSETS / "akla-logo.png"
        if logo.exists() and not self.m.get("no_logo"):
            rp.add_run().add_picture(str(logo), height=Inches(0.55))
        for hp in hdr.paragraphs:
            if not hp.text.strip() and not hp.runs:
                hp._p.getparent().remove(hp._p)
        ftr = sec.footer
        ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style(fp.add_run("Page "), size=8)
        _field(fp, "PAGE")
        _style(fp.add_run(" of "), size=8)
        _field(fp, "NUMPAGES")

    # ------------------------------------------------------------ primitives
    def para(self, text="", align="justify", size=11, color=None, bold=None, italic=None, caps=None,
             left=0.0, first=0.0, after=0, container=None):
        target = container if container is not None else self.doc
        p = target.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        pf = p.paragraph_format
        pf.left_indent = Inches(left)
        pf.first_line_indent = Inches(first)
        pf.space_after = Pt(after)
        if text:
            _rich(p, text, size=size, color=color, bold=bold, italic=italic, caps=caps)
        return p

    def numbered(self, number: str, text: str, container=None, size=11, color=None):
        """Word supplies the number; `number` only fixes the level."""
        lvl = max(0, len(str(number).strip(".").split(".")) - 1)
        p = self.para("", left=0.5, first=-0.5, container=container, size=size)
        _set_num(p, _akla_numbering(self.doc), lvl)
        _rich(p, text, size=size, color=color)
        return p

    def blank(self, container=None, size=11):
        """An empty line. The house rule spaces with these, never with space_after."""
        return self.para("", container=container, size=size)

    def heading(self, number: str, title: str, level=1, container=None):
        target = container if container is not None else self.doc
        p = target.add_paragraph(style=f"Heading {level}")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        _set_num(p, _akla_numbering(self.doc), level - 1)
        _style(p.add_run(title), size=11, bold=(level == 1), caps=True,
               underline=True, color=BLACK)
        if level == 1 and container is None:
            self.landmark(p, f"{str(number).rstrip('.')}. {title}", f"Section{str(number).rstrip('.').replace('.', '_')}")
        return p

    def landmark(self, p, text, name):
        """Bookmark a landmark paragraph and enter it in the contents."""
        self.bookmarks.append((text, name))
        _bookmark(p, len(self.bookmarks), name)
        _tc(p, text)
        return p

    def marked(self, kind: str, text: str, container=None, indent=0.5, label=None):
        col = {"comment": BLUE, "action": RED, "response": GREEN}[kind]
        lab = label or {"comment": "AKLA Comment", "action": "AKLA Action Item", "response": "AKLA Response"}[kind]
        return self.para(f"{lab}: {text}", color=col, left=(0.0 if container is not None else indent),
                         container=container)

    def list_items(self, items, style="alpha", base_left=0.5, container=None, color=None, size=11):
        """One Word list per lettered list: Word supplies (a), (b) ... and restarts at (a) for each list."""
        nid = _new_list(self.doc, style)
        for it in items:
            p = self.para("", left=base_left + 0.5, first=-0.5, container=container)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(base_left + 0.5))
            _set_num(p, nid, 0)
            _rich(p, str(it), size=size, color=color)

    def rest_blank(self):
        self.para("[***THE REST OF THIS PAGE HAS BEEN INTENTIONALLY LEFT BLANK***]",
                  align="center", bold=True, caps=True, after=0)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def centered(self, text, size=12, rule=False, color=BLACK, after=0):
        p = self.para(text, align="center", size=size, bold=True, caps=True, color=color, after=after)
        if rule:
            _bottom_rule(p)
        return p

    # ------------------------------------------------------------ tables
    def table(self, columns, rows, col_widths=None, font_size=10, align=None, container=None):
        target = container if container is not None else self.doc
        t = target.add_table(rows=1, cols=len(columns))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        w = list(col_widths) if col_widths and len(col_widths) == len(columns) else [self.width / len(columns)] * len(columns)
        if sum(w) > self.width or sum(w) < self.width * 0.6:
            k = self.width / sum(w)
            w = [x * k for x in w]
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, c in enumerate(columns):
            cell = hdr.cells[i]
            _shade(cell, NAVY); _cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            _style(p.add_run(c), size=font_size, bold=True, caps=True, color=GOLD)
        for row in rows:
            cells = t.add_row().cells
            if len(row) == 1 and isinstance(row[0], dict) and row[0].get("type") == "group":
                c = cells[0].merge(cells[-1]) if len(cells) > 1 else cells[0]
                _shade(c, GREY); _cell_margins(c)
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                _style(p.add_run(row[0].get("text", "")), size=font_size, bold=True, caps=True)
                continue
            for i, val in enumerate(row):
                if i >= len(cells):
                    break
                _cell_margins(cells[i])
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
                self._fill(cells[i], val, font_size, (align or {}).get(i, "left"))
        _widths(t, w)
        self.blank()
        return t

    def _fill(self, cell, val, font_size=10, align="left"):
        blocks = val if isinstance(val, list) else [val]
        first = True
        for b in blocks:
            if isinstance(b, dict):
                self.block(b, container=cell, size=font_size, first_para=first)
                first = False
                continue
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
            p.paragraph_format.space_after = Pt(0)
            _rich(p, "" if b is None else str(b), size=font_size)
            first = False
        if isinstance(val, list) and val and isinstance(val[0], dict) and len(cell.paragraphs) > 1 and not cell.paragraphs[0].text:
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    def contract_table(self, title: str, groups, label_width=1.9, font_size=10.5):
        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        row = t.add_row()
        _repeat_header(row)
        c = row.cells[0].merge(row.cells[1])
        _shade(c, NAVY); _cell_margins(c)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        _style(p.add_run(title.upper()), size=font_size, bold=True, caps=True, color=GOLD)
        for g in groups:
            if g.get("title"):
                row = t.add_row()
                c = row.cells[0].merge(row.cells[1])
                _shade(c, GREY); _cell_margins(c)
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                _style(p.add_run(g["title"].upper()), size=font_size, bold=True, caps=True)
            for r in g.get("rows", []):
                row = t.add_row()
                lc, rc = row.cells
                _cell_margins(lc); _cell_margins(rc)
                lp = lc.paragraphs[0]; lp.paragraph_format.space_after = Pt(0)
                _rich(lp, r.get("label", ""), size=font_size)
                blocks = r.get("blocks") or [{"type": "para", "text": r.get("text", "")}]
                for i, b in enumerate(blocks):
                    self.block(b, container=rc, size=font_size, first_para=(i == 0))
                if rc.paragraphs and not rc.paragraphs[0].text and len(rc.paragraphs) > 1:
                    rc.paragraphs[0]._p.getparent().remove(rc.paragraphs[0]._p)
        label_width = min(label_width, self.width * 0.3)       # portrait: the label must not crowd the analysis
        _widths(t, [label_width, self.width - label_width])
        self.blank()
        return t

    # ------------------------------------------------------------ block router
    def block(self, b: dict, container=None, size=11, first_para=False):
        t = b.get("type", "para")
        in_cell = container is not None
        if t == "heading":
            self.heading(b.get("number", ""), b.get("text", ""), 1, container)
        elif t == "subheading":
            self.heading(b.get("number", ""), b.get("text", ""), 2, container)
        elif t == "para":
            if b.get("number") and not in_cell:
                self.numbered(b["number"], b.get("text", ""), size=size)
            else:
                p = container.paragraphs[0] if (in_cell and first_para and not container.paragraphs[0].text) else None
                if p is None:
                    p = self.para("", container=container, size=size, left=0.0 if in_cell else b.get("indent", 0.5))
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(0)
                _rich(p, b.get("text", ""), size=size)
        elif t in ("comment", "action", "response"):
            self.marked(t, b.get("text", ""), container=container, label=b.get("label"))
        elif t == "list":
            self.list_items(b.get("items", []), b.get("style", "alpha"),
                            0.0 if in_cell else b.get("indent", 0.5), container, b.get("color"))
        elif t == "table":
            self.table(b.get("columns", []), b.get("rows", []), b.get("col_widths"),
                       b.get("font_size", 10), container=container)
        elif t == "contract_table":
            self.contract_table(b.get("title", ""), b.get("groups", []),
                                min(2.6, max(1.4, b.get("label_width") or 1.9)))
        elif t == "note":
            self.para(b.get("text", ""), size=9, italic=True, left=0.5, container=container)
        elif t == "blank_rest":
            self.rest_blank()
        elif t == "page_break":
            self.page_break()
        elif t == "title":
            self.centered(b.get("text", ""), b.get("size", 12), b.get("rule", False))
        else:
            self.para(b.get("text", ""), container=container, size=size)

    # ------------------------------------------------------------ front/back matter
    def cover(self):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        _shade(c, NAVY); _cell_margins(c, 160, 160, 200, 200)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style(p.add_run(self.m["report_title"]), size=15, bold=True, caps=True, color=GOLD)
        _widths(t, [self.width])
        for _ in range(3):
            self.blank()
        self.para("DATED:", align="right", bold=True, caps=True)
        self.blank()
        self.para(self.m.get("date", ""), align="right", caps=True)
        for _ in range(2):
            self.blank()
        self.para("PREPARED BY:", align="left", bold=True, caps=True)
        self.blank()
        wide = ASSETS / "akla-logo-wide.jpeg"
        if wide.exists() and not self.m.get("no_logo"):
            self.doc.add_paragraph().add_run().add_picture(str(wide), height=Inches(0.5))
        for _ in range(3):
            self.blank()
        self.centered("RANKED BY", size=11, rule=True)
        self.blank()
        rb = ASSETS / "ranked-by.png"
        if rb.exists() and not self.m.get("no_logo"):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(rb), width=Inches(min(5.6, self.width)))
        self.page_break()

    def toc(self, entries):
        """entries: [(text, bookmark)]. A real TOC field over the TC landmarks, with the rows cached (dotted
        leader and PAGEREF page number) so the contents reads correctly before Word refreshes it."""
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        self.centered("TABLE OF CONTENTS", size=11, rule=True)
        self.blank()
        p = self.doc.add_paragraph()
        x = OxmlElement("w:fldChar"); x.set(qn("w:fldCharType"), "begin"); p.add_run()._r.append(x)
        x = OxmlElement("w:instrText"); x.set(qn("xml:space"), "preserve"); x.text = " TOC \\f C \\h \\z "
        p.add_run()._r.append(x)
        x = OxmlElement("w:fldChar"); x.set(qn("w:fldCharType"), "separate"); p.add_run()._r.append(x)
        for text, name in entries:
            p = self.para("", align="left")
            p.paragraph_format.tab_stops.add_tab_stop(Inches(self.width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            _style(p.add_run(text + "\t"), size=10, bold=True, caps=True)
            _field(p, "PAGEREF " + name + " \\h", placeholder="—")
            self.blank()
        p = self.doc.add_paragraph()
        x = OxmlElement("w:fldChar"); x.set(qn("w:fldCharType"), "end"); p.add_run()._r.append(x)
        self.rest_blank()
        self.page_break()

    def important_note(self, blocks):
        p = self.centered("IMPORTANT NOTE", size=11)
        self.landmark(p, "IMPORTANT NOTE", "ImportantNote")
        self.blank()
        for b in blocks:
            if b.get("type") == "list":
                self.list_items(b.get("items", []), b.get("style", "alpha"), base_left=0.0)
            elif b.get("type") == "title":
                self.centered(b.get("text", ""), size=min(12, b.get("size", 10)))
            elif b.get("type") != "blank_rest":
                self.para(b.get("text", ""))
            else:
                continue
            self.blank()
        self.rest_blank()
        self.page_break()

    def glossary(self, entries):
        if not entries:
            return
        p = self.centered("GLOSSARY", size=10, rule=True)
        self.landmark(p, "GLOSSARY", "Glossary")
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, h in enumerate(("DEFINED TERM", "MEANING")):
            c = hdr.cells[i]
            _shade(c, GREY); _cell_margins(c)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style(p.add_run(h), size=10, bold=True, caps=True)
        for e in sorted(entries, key=lambda x: x["term"].lower()):
            row = t.add_row()
            cs = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cs)   # short rows only
            lc, rc = row.cells
            _shade(lc, GREY); _cell_margins(lc); _cell_margins(rc)
            _style(lc.paragraphs[0].add_run(e["term"]), size=10, bold=True)
            rc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _rich(rc.paragraphs[0], e["meaning"], size=10)
        # proportional, not a fixed 2.9in: on a portrait page a fixed term column takes 43% of the width
        _widths(t, [self.width * 0.29, self.width * 0.71])
        self.blank()
        self.rest_blank()
        self.page_break()

    def volume_title_page(self, text):
        for _ in range(8):
            self.blank()
        self.centered(text, size=12)
        self.page_break()
        p = self.centered(text, size=12, rule=True)
        self.landmark(p, text, "VolumeBody")

    def annexure(self, ann):
        if not ann or not ann.get("items"):
            return
        self.page_break()
        title = ann.get("title", "ANNEXURE A – REVIEWED DOCUMENTS")
        p = self.centered(title, size=10, rule=True)
        self.landmark(p, title, "AnnexureA")
        self.blank()
        intro = ann.get("intro") or f"For the purposes of this Volume {self.m['volume_roman']}, the Firm has reviewed the following documents:"
        p = self.para("", left=0.5, first=-0.5)
        _set_num(p, _new_multilevel(self.doc, "AKLA Annexure Numbering"), 0)        # restarts at 1.
        _rich(p, intro)
        self.blank()
        items = [re.sub(r"(;\s*and|[;.])\s*$", "", str(i).strip()).strip() for i in ann["items"]]
        out = [x + (";" if i < len(items) - 2 else ("; and" if i == len(items) - 2 else ".")) for i, x in enumerate(items)]
        self.list_items(out, base_left=0.5)
        self.blank()
        self.rest_blank()

    def numbered_section(self, title, bookmark, blocks, new_page=True):
        """A landmarked section with its own Word numbering (1., 1.1.), so it never shifts the subject-section
        numbers: Scope and Methodology, Executive Summary, Transaction Actions, Annexure B Review Coverage."""
        if not blocks:
            return
        if new_page:
            self.page_break()
        self.landmark(self.centered(title, size=11, rule=True), title, bookmark)
        self.blank()
        nid = _new_multilevel(self.doc, "AKLA " + bookmark + " Numbering")
        under_heading = False
        for b in blocks:
            if b.get("type") == "heading":
                p = self.doc.add_paragraph(style="Heading 1")
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
                _set_num(p, nid, 0)
                _style(p.add_run(b.get("text") or b.get("title", "")), size=11, bold=True, caps=True, underline=True, color=BLACK)
                under_heading = True
            elif b.get("type", "para") == "para" and under_heading:
                p = self.para("", left=0.5, first=-0.5)
                _set_num(p, nid, 1)
                _rich(p, b.get("text", ""))
            else:
                self.block(b)
            self.blank()

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        carry = False                          # a spacer after a heading travels with the heading and its text
        for p in self.doc.paragraphs:
            if p.text.strip():
                carry = p.style.name.startswith("Heading")
            elif carry:
                p.paragraph_format.keep_with_next = True
        self.doc.save(str(path))
        return path


def build(vol: dict, out: Path) -> Path:
    v = Volume(vol)
    if not vol.get("no_cover"):
        v.cover()
    entries = []
    coverage_title = vol.get("coverage_title") or COVERAGE_TITLE
    if vol.get("important_note"):
        entries.append(("IMPORTANT NOTE", "ImportantNote"))
    if vol.get("methodology"):
        entries.append((METHODOLOGY_TITLE, "Methodology"))
    if vol.get("executive_summary"):
        entries.append((SUMMARY_TITLE, "ExecutiveSummary"))
    if vol.get("glossary"):
        entries.append(("GLOSSARY", "Glossary"))
    title_page = vol.get("volume_title_page") or f"VOLUME {vol['volume_roman']} – {vol['volume_title'].upper()}"
    entries.append((title_page, "VolumeBody"))
    entries += [(f"{b['number'].rstrip('.')}.  {b['text']}", f"Section{b['number'].rstrip('.').replace('.', '_')}")
                for b in vol.get("blocks", []) if b.get("type") == "heading"]
    if vol.get("transaction_actions"):
        entries.append((ACTIONS_TITLE, "TransactionActions"))
    if vol.get("annexure", {}).get("items"):
        entries.append((vol["annexure"].get("title", "ANNEXURE A – REVIEWED DOCUMENTS"), "AnnexureA"))
    if vol.get("coverage"):
        entries.append((coverage_title, "Coverage"))
    v.toc(entries)
    if vol.get("important_note"):
        v.important_note(vol["important_note"])
    # optional sections, each numbered independently: the Important Note already ends on a page break
    v.numbered_section(METHODOLOGY_TITLE, "Methodology", vol.get("methodology"), new_page=False)
    v.numbered_section(SUMMARY_TITLE, "ExecutiveSummary", vol.get("executive_summary"), new_page=bool(vol.get("methodology")))
    if vol.get("methodology") or vol.get("executive_summary"):
        v.page_break()
    v.glossary(vol.get("glossary") or [])
    v.volume_title_page(title_page)
    for i, b in enumerate(vol.get("blocks", [])):
        if i and b.get("type") not in ("page_break", "blank_rest"):
            v.blank()                              # the gap is an explicit line, not paragraph spacing
        if b.get("type") != "blank_rest":
            v.block(b)
    v.blank()
    v.rest_blank()
    v.numbered_section(ACTIONS_TITLE, "TransactionActions", vol.get("transaction_actions"))
    v.annexure(vol.get("annexure"))
    v.numbered_section(coverage_title, "Coverage", vol.get("coverage"))
    sec = v.doc.sections[0]
    for part in (v.doc.part, sec.header.part, sec.footer.part):
        _zero_spacing(part.element)
    for st in v.doc.styles:
        if st.type == 1:
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(0)
    return v.save(out)


def main():
    ap = argparse.ArgumentParser(description="Render a DD volume JSON into AKLA gold-standard Word format.")
    ap.add_argument("json_file")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--draft-label", help="override, e.g. 'SECOND DRAFT'")
    ap.add_argument("--date", help="override, e.g. 'September 07, 2026'")
    ap.add_argument("--no-logo", action="store_true")
    ap.add_argument("--landscape", action="store_true",
                    help="exception only: a wide schedule that cannot be restructured, or a re-issue of a volume circulated in landscape")
    a = ap.parse_args()
    vol = json.loads(Path(a.json_file).read_text())
    if a.draft_label:
        vol["draft_label"] = a.draft_label
    if a.date:
        vol["date"] = a.date
    if a.no_logo:
        vol["no_logo"] = True
    if a.landscape:
        vol["orientation"] = "landscape"
    p = build(vol, Path(a.out))
    n_marks = sum(1 for b in vol.get("blocks", []) if b.get("type") in ("comment", "action", "response"))
    print(f"written: {p}")
    print(f"  {len(vol.get('blocks', []))} body blocks, {len(vol.get('glossary') or [])} glossary terms, "
          f"{n_marks} top-level comments/action items")


if __name__ == "__main__":
    main()
