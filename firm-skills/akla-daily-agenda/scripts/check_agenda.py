#!/usr/bin/env python3
"""Check a built Daily Agenda .docx against the house format and print it as text.

    python3 check_agenda.py "Daily Agenda - Name [October 07, 2026].docx" [--render DIR]

Exit code 1 on any FAIL. --render writes page PNGs via LibreOffice (if installed)
so the result can be looked at before it is sent.
"""
import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml.ns import qn

LABELS = ["Transaction", "Scope of Task", "Complete Deliverable", "Day Deliverable", "Current Deliverable Status"]
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--render", metavar="DIR", help="write page PNGs here (needs LibreOffice + pymupdf)")
    a = ap.parse_args()
    doc = Document(a.docx)
    fails, warns = [], []

    # --- geometry
    s = doc.sections[0]
    if abs(s.page_width.inches - 8.27) > 0.05 or abs(s.page_height.inches - 11.69) > 0.05:
        fails.append(f"page is not A4 portrait ({s.page_width.inches:.2f} x {s.page_height.inches:.2f} in)")

    # --- text + fonts
    paras = list(doc.paragraphs)
    texts = [p.text for p in paras]
    body = "\n".join(texts)
    for p in paras:
        for r in p.runs:
            if r.text.strip() and r.font.name not in (None, "Arial"):
                fails.append(f"non-Arial run: {r.font.name!r} in {p.text[:40]!r}")
                break
    for must in ("Daily Agenda", "Daily Timesheet", "Deliverables"):
        if must not in body:
            fails.append(f"missing '{must}'")

    # --- dates
    dates = re.findall(r"Date: ([A-Z][a-z]+ \d{2}, \d{4})", body)
    if len(dates) != 2:
        fails.append(f"expected two 'Date:' lines (agenda, timesheet), found {len(dates)}")
    else:
        try:
            ag, ts = (dt.datetime.strptime(x, "%B %d, %Y").date() for x in dates)
            if ag <= ts:
                fails.append(f"agenda date {ag} is not after timesheet date {ts}")
            if ag.weekday() >= 5:
                warns.append(f"agenda date {ag} falls on a weekend")
        except ValueError as e:
            fails.append(f"date format: {e}")

    # --- deliverable blocks
    n_tx = texts.count if False else sum(1 for t in texts if t.startswith("Transaction: "))
    for lab in LABELS:
        n = sum(1 for t in texts if t.startswith(lab + ": "))
        if n != n_tx:
            fails.append(f"'{lab}' appears {n} times but there are {n_tx} transactions")
    for t in texts:
        if t.startswith(tuple(l + ": " for l in LABELS[1:])) and not t.rstrip().endswith((".", "?", "!")):
            warns.append(f"no full stop: {t[:70]!r}")
        if t.startswith("Scope of Task: ") and len(t) < 25:
            warns.append(f"very short scope: {t!r}")

    # --- table
    if not doc.tables:
        fails.append("no timesheet table")
    else:
        t = doc.tables[0]
        hdr = [c.text.replace("\n", " ") for c in t.rows[0].cells]
        for h in ("Sr. No.", "Transaction", "Matter & Description", "Hours", "Billable Hours"):
            if h not in hdr:
                fails.append(f"table header missing '{h}' (got {hdr})")
        for c in t.rows[0].cells:
            shd = c._tc.tcPr.find(qn("w:shd")) if c._tc.tcPr is not None else None
            if shd is None or shd.get(qn("w:fill")) != "002060":
                fails.append("header row is not navy (002060)")
                break
            for p in c.paragraphs:
                for r in p.runs:
                    if r.text.strip() and (not r.bold or str(r.font.color.rgb) != "FFC000"):
                        fails.append("header text is not bold gold (FFC000)")
                        break
        hi = hdr.index("Hours") if "Hours" in hdr else 3
        data = [r for r in t.rows[1:] if not r.cells[0].text.startswith(("Total", "Percentage"))]
        total_row = next((r for r in t.rows if r.cells[0].text.startswith("Total")), None)

        def num(x):
            x = x.strip()
            return 0.0 if x in ("", "-") else float(x)

        try:
            summed = sum(num(r.cells[hi].text) for r in data)
            if total_row is None:
                fails.append("no Total Hours row")
            elif abs(num(total_row.cells[hi].text) - summed) > 0.01:
                fails.append(f"Total Hours {total_row.cells[hi].text} != sum of rows {summed:g}")
            if summed > 14:
                warns.append(f"{summed:g} hours in one day — confirm with the author")
            if summed < 4:
                warns.append(f"only {summed:g} hours recorded — confirm nothing is missing")
        except ValueError as e:
            fails.append(f"non-numeric hours in table: {e}")
        for r in data:
            if not r.cells[2].text.strip():
                fails.append("timesheet row with empty description")
            if not r.cells[0].text.strip():
                warns.append("timesheet row without a serial number")

    # --- optional render
    pages = None
    if a.render or os.path.exists(SOFFICE):
        with tempfile.TemporaryDirectory() as tmp:
            try:
                subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", tmp, a.docx],
                               check=True, capture_output=True, timeout=120)
                pdf = os.path.join(tmp, os.path.splitext(os.path.basename(a.docx))[0] + ".pdf")
                try:
                    import pymupdf
                    d = pymupdf.open(pdf)
                    pages = d.page_count
                    if a.render:
                        os.makedirs(a.render, exist_ok=True)
                        for i, pg in enumerate(d, start=1):
                            out = os.path.join(a.render, f"page{i}.png")
                            pg.get_pixmap(dpi=110).save(out)
                            print(f"rendered {out}")
                except ImportError:
                    warns.append("pymupdf not installed: page count / render skipped")
            except Exception as e:  # noqa: BLE001
                warns.append(f"LibreOffice render failed: {e}")
    if pages is not None and pages != 2:
        warns.append(f"document runs to {pages} pages (the firm's agendas are two: agenda, timesheet)")

    # --- report
    print("=" * 70)
    print(body)
    if doc.tables:
        print("-" * 70)
        for r in doc.tables[0].rows:
            print(" | ".join(c.text.replace("\n", " / ") for c in r.cells))
    print("=" * 70)
    for w in warns:
        print("WARN", w)
    for f in fails:
        print("FAIL", f)
    if not fails:
        print(f"OK   {os.path.basename(a.docx)}" + (f" ({pages} pages)" if pages else ""))
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
