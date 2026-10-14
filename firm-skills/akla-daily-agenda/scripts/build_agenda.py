#!/usr/bin/env python3
"""Build an AKLA Daily Agenda + Daily Timesheet .docx from a YAML spec.

    python3 build_agenda.py agenda.yaml [-o out.docx] [--outdir DIR]

The spec is documented in references/agenda-format.md and examples/example.yaml.
A per-person profile at ~/.akla-agenda.yaml (name, role, closing, ak_column) is
read first; keys in the spec override it.
"""
import argparse
import datetime as dt
import os
import re
import sys

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Twips

NAVY = "002060"
GOLD = RGBColor(0xFF, 0xC0, 0x00)
FONT = "Arial"
BODY_PT = 11
TITLE_PT = 13
PROFILE = os.path.expanduser("~/.akla-agenda.yaml")

CLOSING_ALLOCATION = "[Subject To Allocation Of Tasks During The Day]"
CLOSING_BLANK = "[***The Rest Of This Page Has Been Intentionally Left Blank***]"
DEFAULT_NOTE = "To be discussed with Boss."

# column widths in twips: page is A4 with 1" margins = 9029 twips usable
COLS_AK = [("Sr. No.", 800), ("Transaction", 1800), ("Matter &\nDescription", 3250),
           ("Hours", 900), ("Billable Hours", 1150), ("AK\nBillable Hours", 1130)]
COLS_NO_AK = [("Sr. No.", 800), ("Transaction", 1800), ("Matter &\nDescription", 4380),
              ("Hours", 900), ("Billable Hours", 1150)]


# ----------------------------------------------------------------- helpers
def die(msg):
    sys.exit(f"build_agenda: {msg}")


def parse_date(v, label):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d %B %Y", "%B %d, %Y", "%B %d %Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    die(f"{label}: cannot read date {v!r} (use YYYY-MM-DD)")


def next_working_day(d):
    d = d + dt.timedelta(days=1)
    while d.weekday() >= 5:  # Sat=5, Sun=6
        d += dt.timedelta(days=1)
    return d


def fmt_date(d):
    return d.strftime("%B %d, %Y")  # "October 07, 2026"


def fmt_hours(x):
    if x is None:
        return "-"
    x = float(x)
    if x == 0:
        return "-"
    return f"{x:.1f}" if abs(x * 10 - round(x * 10)) < 1e-9 else f"{x:.2f}"


def to_float(v, label):
    if v in (None, "", "-"):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        die(f"{label}: hours must be a number, got {v!r}")


def ensure_period(s):
    s = str(s).strip()
    if s and s[-1] not in ".?!":
        s += "."
    return s


# ---------------------------------------------------------------- styling
def set_run_font(run, bold=None, italic=None, underline=None, size=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size or BODY_PT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if color is not None:
        run.font.color.rgb = color


def style_normal(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.07
    for name in ("List Bullet", "List Paragraph"):
        try:
            doc.styles[name].font.name = FONT
            doc.styles[name].font.size = Pt(BODY_PT)
        except KeyError:
            pass


def para(doc_or_cell, align=None, style=None):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    return p


def add_runs(p, parts):
    """parts: list of (text, dict(bold=,italic=,underline=,size=))"""
    for text, kw in parts:
        r = p.add_run(text)
        set_run_font(r, **kw)
    return p


def label_line(doc, label, value, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = para(doc, align)
    add_runs(p, [(label, dict(bold=True)), (": " + value, dict(bold=False))])
    return p


def heading_line(doc, text):
    p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_runs(p, [(text, dict(bold=True, underline=True))])
    return p


def title_block(doc, title, date, role, name):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(p, [(title, dict(bold=True, size=TITLE_PT))])
    para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    label_line(doc, "Date", fmt_date(date), WD_ALIGN_PARAGRAPH.RIGHT)
    label_line(doc, role, name, WD_ALIGN_PARAGRAPH.RIGHT)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def valign_center(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)


def set_cell_width(cell, twips):
    cell.width = Twips(twips)


def table_margins(table, twips=115):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def set_grid(table, widths):
    """Write the column widths into tblGrid + fixed layout so Word and LibreOffice honour them."""
    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def cell_text(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, color=None):
    """Write text into a cell. '\n\n' separates paragraphs with a blank line between."""
    cell.text = ""
    first = True
    chunks = str(text).split("\n\n")
    for i, chunk in enumerate(chunks):
        if i > 0:
            para(cell)  # blank line between items, as the firm does
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = align
        lines = chunk.split("\n")
        for j, line in enumerate(lines):
            r = p.add_run(line)
            set_run_font(r, bold=bold, color=color)
            if j < len(lines) - 1:
                r.add_break()
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)


def join_items(items):
    """['A', 'B', 'C'] -> 'A;\n\nB; and\n\nC.' — the firm's list-in-a-cell style."""
    items = [str(i).strip().rstrip(".;") for i in items if str(i).strip()]
    if not items:
        return ""
    if len(items) == 1:
        return items[0] + "."
    out = []
    for i, it in enumerate(items):
        if i == len(items) - 1:
            out.append(it + ".")
        elif i == len(items) - 2:
            out.append(it + "; and")
        else:
            out.append(it + ";")
    return "\n\n".join(out)


# ------------------------------------------------------------------- build
def build(spec, out_path):
    name = spec.get("name") or die("name is required (or set it in ~/.akla-agenda.yaml)")
    role = spec.get("role") or die("role is required, e.g. Associate / Senior Counsel / Intern")
    ts_date = parse_date(spec.get("timesheet_date"), "timesheet_date") or dt.date.today()
    ag_date = parse_date(spec.get("agenda_date"), "agenda_date") or next_working_day(ts_date)
    if ag_date <= ts_date:
        die(f"agenda_date {ag_date} must be after timesheet_date {ts_date}")
    note = spec.get("note", DEFAULT_NOTE)
    closing = (spec.get("closing") or "allocation").lower()
    if closing not in ("allocation", "blank"):
        die("closing must be 'allocation' or 'blank'")
    ak_column = bool(spec.get("ak_column", True))
    deliverables = spec.get("deliverables") or []
    internal = spec.get("internal") or []
    timesheet = spec.get("timesheet") or []
    if not deliverables and not internal:
        die("no deliverables and no internal items: nothing to put on the agenda")
    if not timesheet:
        die("timesheet is empty: nothing to record for today")

    doc = Document()
    style_normal(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Inches(1))

    # ------------------------------------------------ page 1: Daily Agenda
    title_block(doc, "Daily Agenda", ag_date, role, name)
    para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    heading_line(doc, "Deliverables ")
    if note:
        p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, style="List Bullet")
        add_runs(p, [("*", dict(bold=True)), ("Note: " + ensure_period(note), dict(italic=True))])
    para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)

    for i, d in enumerate(deliverables):
        tx = d.get("transaction") or die(f"deliverables[{i}]: transaction is required")
        scope = d.get("scope") or die(f"deliverables[{i}] ({tx}): scope is required")
        complete = d.get("complete_deliverable") or "N/A"
        day = d.get("day_deliverable") or complete
        status = d.get("status") or "Pending with me."
        label_line(doc, "Transaction", str(tx).strip())
        label_line(doc, "Scope of Task", ensure_period(scope))
        label_line(doc, "Complete Deliverable", ensure_period(complete))
        label_line(doc, "Day Deliverable", ensure_period(day))
        label_line(doc, "Current Deliverable Status", ensure_period(status))
        para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)

    if internal:
        p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_runs(p, [("Internal/Administrative", dict(bold=True, underline=True)), (":", dict())])
        para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
        for item in internal:
            p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, style="List Bullet")
            add_runs(p, [(ensure_period(item), dict())])
        para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(p, [(CLOSING_ALLOCATION if closing == "allocation" else CLOSING_BLANK, dict())])
    p.add_run().add_break(WD_BREAK.PAGE)

    # --------------------------------------------- page 2: Daily Timesheet
    title_block(doc, "Daily Timesheet", ts_date, role, name)
    para(doc)

    cols = COLS_AK if ak_column else COLS_NO_AK
    ncols = len(cols)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_margins(table)
    set_grid(table, [w for _, w in cols])
    hdr = table.rows[0]
    repeat_header(hdr)
    for c, (title, w) in enumerate(cols):
        cell = hdr.cells[c]
        set_cell_width(cell, w)
        shade(cell, NAVY)
        valign_center(cell)
        cell_text(cell, title, WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=GOLD)

    tot = tot_b = tot_ak = 0.0
    for i, row in enumerate(timesheet, start=1):
        tx = row.get("transaction") or die(f"timesheet[{i}]: transaction is required")
        desc = row.get("description") or die(f"timesheet[{i}] ({tx}): description is required")
        if isinstance(desc, (list, tuple)):
            desc = join_items(desc)
        else:
            desc = ensure_period(desc)
        if "hours" not in row or row["hours"] in (None, ""):
            die(f"timesheet[{i}] ({tx}): hours is required — ask, never guess")
        h = to_float(row["hours"], f"timesheet[{i}].hours")
        b = to_float(row.get("billable"), f"timesheet[{i}].billable")
        ak = to_float(row.get("ak_billable"), f"timesheet[{i}].ak_billable")
        if b > h or ak > h:
            die(f"timesheet[{i}] ({tx}): billable hours exceed hours worked")
        tot, tot_b, tot_ak = tot + h, tot_b + b, tot_ak + ak
        r = table.add_row()
        vals = [f"{i}.", str(tx).strip(), desc, fmt_hours(h), fmt_hours(b)] + ([fmt_hours(ak)] if ak_column else [])
        for c, v in enumerate(vals):
            set_cell_width(r.cells[c], cols[c][1])
            align = WD_ALIGN_PARAGRAPH.JUSTIFY if c == 2 else WD_ALIGN_PARAGRAPH.CENTER
            if c == 1:
                align = WD_ALIGN_PARAGRAPH.LEFT
            cell_text(r.cells[c], v, align)
            if c != 2:
                valign_center(r.cells[c])

    def summary_row(label, values):
        r = table.add_row()
        merged = r.cells[0].merge(r.cells[2])
        cell_text(merged, label, WD_ALIGN_PARAGRAPH.CENTER)
        for c, v in enumerate(values, start=3):
            set_cell_width(r.cells[c], cols[c][1])
            cell_text(r.cells[c], v, WD_ALIGN_PARAGRAPH.CENTER)
            valign_center(r.cells[c])

    summary_row("Total Hours", [fmt_hours(tot), fmt_hours(tot_b)] + ([fmt_hours(tot_ak)] if ak_column else []))

    def pct(x):
        return f"{x / tot * 100:.0f}%" if tot and x else "-"

    summary_row("Percentage of Billable Hours", ["-", pct(tot_b)] + ([pct(tot_ak)] if ak_column else []))

    para(doc)
    if closing == "blank":
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(p, [(CLOSING_BLANK, dict())])

    doc.save(out_path)
    return dict(name=name, role=role, agenda_date=ag_date, timesheet_date=ts_date,
                total=tot, billable=tot_b, ak=tot_ak, rows=len(timesheet),
                deliverables=len(deliverables), internal=len(internal))


def default_filename(name, ag_date):
    safe = re.sub(r"[\\/:*?\"<>|]", "", name).strip()
    return f"Daily Agenda - {safe} [{fmt_date(ag_date)}].docx"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="YAML spec (see examples/example.yaml)")
    ap.add_argument("-o", "--out", help="output .docx (default: house file name)")
    ap.add_argument("--outdir", default=".", help="folder for the default file name")
    a = ap.parse_args()

    spec = {}
    if os.path.exists(PROFILE):
        with open(PROFILE) as f:
            spec.update(yaml.safe_load(f) or {})
    with open(a.spec) as f:
        spec.update(yaml.safe_load(f) or {})

    ts = parse_date(spec.get("timesheet_date"), "timesheet_date") or dt.date.today()
    ag = parse_date(spec.get("agenda_date"), "agenda_date") or next_working_day(ts)
    out = a.out or os.path.join(a.outdir, default_filename(spec.get("name", "Unnamed"), ag))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    info = build(spec, out)
    print(f"Wrote {out}")
    print(f"  {info['role']}: {info['name']}")
    print(f"  Agenda for {fmt_date(info['agenda_date'])}: {info['deliverables']} transaction(s), {info['internal']} internal item(s)")
    print(f"  Timesheet for {fmt_date(info['timesheet_date'])}: {info['rows']} row(s), "
          f"{fmt_hours(info['total'])} hours, billable {fmt_hours(info['billable'])}, AK billable {fmt_hours(info['ak'])}")


if __name__ == "__main__":
    main()
